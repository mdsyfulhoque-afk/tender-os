from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.section import WD_ORIENT
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from pathlib import Path

OUT = Path('docs/concept-map-tenderly-to-tenderos.docx')

def shade(cell, fill):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement('w:shd'); shd.set(qn('w:fill'), fill); tcPr.append(shd)

def set_cell_text(cell, text, bold=False, color=None, size=8.1):
    cell.text = ''
    p = cell.paragraphs[0]
    p.paragraph_format.space_after = Pt(0)
    p.paragraph_format.space_before = Pt(0)
    r = p.add_run(text); r.bold = bold; r.font.name = 'Aptos'; r.font.size = Pt(size)
    if color: r.font.color.rgb = RGBColor(*color)
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER

def set_repeat_table_header(row):
    trPr = row._tr.get_or_add_trPr(); tblHeader = OxmlElement('w:tblHeader'); tblHeader.set(qn('w:val'), 'true'); trPr.append(tblHeader)
def prevent_row_split(row):
    trPr=row._tr.get_or_add_trPr(); node=OxmlElement('w:cantSplit'); trPr.append(node)

def set_cell_margins(cell, top=65, start=75, bottom=65, end=75):
    tc = cell._tc; tcPr = tc.get_or_add_tcPr(); tcMar = tcPr.first_child_found_in('w:tcMar')
    if tcMar is None: tcMar = OxmlElement('w:tcMar'); tcPr.append(tcMar)
    for m,v in [('top',top),('start',start),('bottom',bottom),('end',end)]:
        node = tcMar.find(qn('w:'+m))
        if node is None: node=OxmlElement('w:'+m); tcMar.append(node)
        node.set(qn('w:w'),str(v)); node.set(qn('w:type'),'dxa')

doc = Document()
sec = doc.sections[0]
sec.orientation = WD_ORIENT.LANDSCAPE
sec.page_width, sec.page_height = Inches(11.7), Inches(8.3)
sec.top_margin = Inches(.55); sec.bottom_margin= Inches(.55); sec.left_margin= Inches(.62); sec.right_margin= Inches(.62)
styles=doc.styles
styles['Normal'].font.name='Aptos'; styles['Normal'].font.size=Pt(9); styles['Normal'].font.color.rgb=RGBColor(44,54,67)
styles['Normal'].paragraph_format.space_after=Pt(5)
for nm,sz,col in [('Title',25,(18,48,72)),('Heading 1',16,(18,82,112)),('Heading 2',11,(18,82,112))]:
    st=styles[nm]; st.font.name='Aptos Display'; st.font.size=Pt(sz); st.font.bold=True; st.font.color.rgb=RGBColor(*col)
# Remove the built-in Title style's decorative paragraph rule.
title_ppr = styles['Title']._element.get_or_add_pPr()
title_bdr = title_ppr.find(qn('w:pBdr'))
if title_bdr is not None: title_ppr.remove(title_bdr)

p=doc.add_paragraph(style='Title'); p.add_run('Tenderly to TenderOS Concept Map')
title_bdr = p._p.get_or_add_pPr().find(qn('w:pBdr'))
if title_bdr is not None: p._p.get_or_add_pPr().remove(title_bdr)
p=doc.add_paragraph(); r=p.add_run('A feature mapping and project development framework'); r.bold=True; r.font.size=Pt(12); r.font.color.rgb=RGBColor(76,103,123)
p=doc.add_paragraph(); p.add_run('Prepared 19 October 2026  |  TenderOS means the TenderProof OS / ProposalGuard local-first diagnostic MVP in this workspace.').italic=True

doc.add_heading('Purpose and design conclusion', level=1)
doc.add_paragraph('Tenderly is a useful product architecture reference because it unifies a working console, scenario simulation, monitoring and response, and infrastructure around one operational record. TenderOS can adapt that product logic to tender decisions: organize each opportunity as a case, model bid scenarios before commitment, inspect source-to-recommendation evidence, monitor deadlines and unresolved conditions, and produce governed outputs. The domains are different: Tenderly manages on-chain systems; TenderOS manages tender/proposal work. The mapping transfers product patterns, not blockchain features.')
doc.add_paragraph('The existing MVP already has a strong case workflow and evidence-centered analysis. The recommended direction is to evolve it into a case operating layer with controlled scenario planning, traceable analysis, deadline/gap monitoring, and a reusable knowledge layer. Treat proposed capabilities as roadmap concepts, not as currently implemented functions.')

# status key
p=doc.add_paragraph(); r=p.add_run('Mapping status:  '); r.bold=True
for label,color in [('Implemented','#DCEFE5'),('Partial','#FFF0CE'),('Proposed','#E4ECF5'),('Not applicable','#ECEFF2')]:
    r=p.add_run(label+'   '); r.font.size=Pt(8); r.font.bold=True
    # status colors shown in table rows, legend intentionally simple

doc.add_heading('Feature by feature mapping', level=1)
doc.add_paragraph('Status reflects the repository as inspected for this concept map. “Partial” means a related workflow or record exists, while the Tenderly pattern still needs a deliberate end-to-end implementation.')
rows = [
('Tenderly Console: shared operational workspace, projects and workspaces','TenderOS Case Workspace: client, tender, donor, country, sector, deadline; case status and coordinated diagnostic pages.','Implemented','TenderOS already centers work on diagnostic cases. Team collaboration and multi-user access are outside the local MVP.'),
('Explorer: unified search across transactions, wallets, contracts and chains','TenderOS Opportunity and Case Explorer: search cases, source files, requirements, evidence, outputs, and reusable IP.','Partial','Case lists and case selection exist; cross-entity search, facets, saved views, and global navigation are roadmap items.'),
('Decoded activity timeline and transaction overview','Case Activity Timeline: source intake, extraction, edits, approvals, score changes, verification and export events.','Partial','AuditEvent exists; a complete human-readable event timeline and object-level provenance view would extend it.'),
('Debugger: step through execution and locate exact failure','Evidence Trace / Decision Inspector: trace each recommendation or readiness flag back through rule, requirement, source page/section, evidence and reviewer.','Proposed','Build explainable, source-linked decisions; do not imply that this is a literal code debugger.'),
('Full error and stack traces with readable messages','Diagnostic Exceptions and Extraction Review: actionable parse/OCR/validation errors with source location, remedy and reviewer disposition.','Partial','Source handling and extraction-failure gap recording exist; consistent error taxonomy and inline remediation are future work.'),
('Custom views, filters, grouping, address labels and tags','Case portfolio views, filters and labels: sector, donor, deadline, readiness, lead status, risk, and owner.','Partial','Case metadata and lead concepts exist; portfolio-level configurable views and tagging are roadmap items.'),
('Shared teams, comments and collaboration','Reviewer assignments, comments, decisions, and controlled hand-offs across proposal, technical, finance and leadership roles.','Proposed','Current application is explicitly local-first/internal with no multi-tenant auth or team collaboration.'),
('Transaction Simulator UI: vary parameters and rerun before committing','Bid Scenario Workbench: compare pursue/no-bid, partner/subconsultant, expert-team, price/budget, schedule and evidence-collection scenarios.','Partial','Bid/no-bid scorecard and recommendation exist; scenario comparison, assumptions and saved variants are not implemented.'),
('Simulation API and reusable simulation calls','Rules and Analysis Service: deterministic scoring, requirement normalization, readiness checks, and scenario calculations callable from UI and future integrations.','Partial','Services exist for scoring, normalization and verification; there is no public API or external AI service connected.'),
('Simulation bundles: sequence multiple actions atomically','Bid Plan / Scenario Bundle: model linked changes together (for example, partner + expert substitution + schedule adjustment) and show combined effects.','Proposed','Treat as a planning workspace with explicit assumptions; never silently change the submitted baseline.'),
('Virtual Environments: fork production state, override and replay history','Case Sandbox / Versioned Working Copy: duplicate a case baseline, test alternative compliance/evidence/methodology plans, compare with the source case.','Proposed','Local case data and deterministic anchors provide a base. Branching, versioning, diff and merge controls are new.'),
('Historical replay from any past block; fork saved simulation','Case Snapshots and Replay: reproduce a prior recommendation from its source versions, scoring rules, evidence and timestamp.','Proposed','Audit events exist, but immutable snapshots of inputs, rule versions and reproducible scoring are not established.'),
('Public Virtual Environment explorer for auditors','Read-only Client/Audit Package: share a scoped, versioned diagnostic with traceable sources and clear readiness labels.','Partial','Reports/workbooks and output package records exist; secure sharing, access scopes and reproducibility are future work.'),
('Transaction traces, state changes and fund-flow visualization','Requirement-to-Evidence Traceability: navigate ToR clause -> normalized requirement -> compliance row -> evidence/gap -> risk -> proposal section -> output.','Partial','Stable match_key/canonical_anchor, compliance rows, gap items, risks and proposal skeleton exist; unified visual trace is a next step.'),
('Gas profiler and performance optimization','Effort, Cost and Delivery Feasibility: expose budget attractiveness, staff loading, deadline feasibility, and effort assumptions in the decision record.','Partial','Scorecard covers budget/value and deadline feasibility; quantified effort/cost models are not documented as implemented.'),
('Contracts and verified source metadata','Tender Source Register: retain ToR/amendment/addendum identity, upload, extracted text, date/version, section anchors and review state.','Partial','Sources and extracted text exist. Amendment comparison, provenance completeness and version conflict handling need development.'),
('Contract read/write interfaces and protocol modules','Tender Rules and Donor Profiles: reusable eligibility/compliance rules, sector methodology modules, templates and known donor instructions.','Partial','Prompt/template/reference modules and reusable IP capture exist; governed rule packs and their lifecycle are future work.'),
('Alerts on events, state changes, balances and function calls','Tender Alerts: deadline milestones, addenda, missing mandatory evidence, unresolved critical risks, stale inputs and readiness regression.','Partial','Workflow guidance, gaps/risks and e-GP stages exist; event-driven notification and escalation are not implemented.'),
('Multichain monitoring with one rule','Portfolio Monitoring: apply shared control checks across all active cases and show exceptions by client, donor, country and deadline.','Proposed','The current MVP is case oriented and has no centralized portfolio monitoring layer.'),
('Web3 Actions: automated code triggered by alert','Controlled Tender Actions: create tasks, draft clarification questions, request missing inputs, or prepare review queues after a trigger.','Partial','Clarification question and workflow optimization concepts exist. Human approval must gate external communications and consequential changes.'),
('Alert routing to email, Slack, PagerDuty, webhooks','Notification Routing: in-app queue first; later authorized email/chat integrations with per-user preferences and delivery audit.','Proposed','No connector or notification routing is present in the local MVP.'),
('Node RPC: production infrastructure, auth, rate limits and analytics','TenderOS Platform Foundation: local database/storage today; future authenticated service/API, backup, access control, integration and operational telemetry.','Partial','SQLite/local storage and local launcher are implemented. SaaS hosting, billing, multi-tenant identity and production SLA are explicitly out of scope.'),
('CI/CD and staging on realistic state; simulation in scripts/bots','Release and Case Verification: deterministic checks for source coverage, mandatory criteria, risk/gap state, prohibited report language, and export readiness.','Implemented','Verification checks and readiness labels exist. Automated deployment pipelines and external integration tests are separate platform work.'),
('Support operations: inspect a case and resolve issues with full context','Diagnostic Delivery Desk: triage blocked inputs, extraction issues, client questions, decision rationale and output revisions.','Partial','Workflow optimizer, audit log, clarification questions and reports support this pattern; service queue/ownership SLAs are not present.'),
('Agentic workflows validated by simulation before execution','Evidence-grounded Assistance: optional AI may draft extraction or analysis; deterministic validation and human review control acceptance.','Proposed','Prompt files are present, but no external AI API or model is connected. AI output must remain a suggestion and cite supplied evidence.'),
('Supported networks, chain-specific capability surface','Tender Context Packs: donor, procurement regime, country, sector, language/OCR and tender portal differences, with explicit applicability.','Partial','Bangladesh e-GP readiness and reference notes exist; broader donor/country packs and coverage controls are future work.'),
]
t=doc.add_table(rows=1, cols=4); t.alignment=WD_TABLE_ALIGNMENT.CENTER; t.autofit=False
widths=[2.45,2.9,0.82,3.65]
for c,w,txt in zip(t.rows[0].cells,widths,['Tenderly capability','TenderOS equivalent concept','Status','TenderOS fit and development note']):
    c.width=Inches(w); set_cell_text(c,txt,True,(255,255,255),8.2); shade(c,'145270'); set_cell_margins(c)
set_repeat_table_header(t.rows[0])
for idx,row in enumerate(rows):
    row_obj=t.add_row(); prevent_row_split(row_obj); cells=row_obj.cells
    for c,w,txt in zip(cells,widths,row): c.width=Inches(w); set_cell_text(c,txt,size=7.55); set_cell_margins(c)
    fill='F2F6F8' if idx%2==0 else 'FFFFFF'
    for c in cells: shade(c,fill)
    st=row[2]; shade(cells[2], {'Implemented':'DCEFE5','Partial':'FFF0CE','Proposed':'E4ECF5','Not applicable':'ECEFF2'}[st]); set_cell_text(cells[2],st,True,(31,71,92),7.2)

# Page break to make the architecture section deliberate
doc.add_page_break()
doc.add_heading('TenderOS layer architecture', level=1)
doc.add_paragraph('The diagram adapts Tenderly’s product architecture pattern while grounding every TenderOS layer in the current MVP or labeling it as future work. The proposal is a layered, case-centered system; it does not imply blockchain infrastructure.')
mermaid = """+--------------------------------------------------------------+
| EXPERIENCE | Case workspace, portfolio, trace, reports      |
+--------------------------------------------------------------+
                              |
                              v
+--------------------------------------------------------------+
| WORKFLOW | Intake > Extract > Assess > Decide > Draft       |
| Verify > Package | Deadlines, tasks, approvals (future)     |
+--------------------------------------------------------------+
                              |
                              v
+--------------------------------------------------------------+
| ANALYSIS | Requirements, compliance, CV, bid score, risks   |
| Scenario workbench (partial / future)                        |
+--------------------------------------------------------------+
                              |
                              v
+--------------------------------------------------------------+
| CASE DOMAIN | Opportunity, source, requirement, evidence,    |
| gap, risk, decision, proposal, output, e-GP checklist        |
+--------------------------------------------------------------+
                              |
                              v
+--------------------------------------------------------------+
| DATA AND KNOWLEDGE | SQLite + files, templates, prompts, IP |
| Versioned snapshots and search index (future)                |
+--------------------------------------------------------------+

  +----------------------------------------------------------+
  | GOVERNANCE CROSS-CUT | Human review, source lineage,      |
  | audit trail and readiness checks; extend version/access  |
  +----------------------------------------------------------+

  +----------------------------------------------------------+
  | PLATFORM FOUNDATION | Local Streamlit, Python services,  |
  | database migrations; backup/API/hosting are future work  |
  +----------------------------------------------------------+"""

for line in mermaid.splitlines():
    p=doc.add_paragraph(); p.paragraph_format.space_after=Pt(0); p.paragraph_format.space_before=Pt(0)
    r=p.add_run(line); r.font.name='Consolas'; r.font.size=Pt(8); r.font.color.rgb=RGBColor(34,64,83)
doc.add_paragraph('Layer flow is shown as an editable ASCII architecture diagram.')

doc.add_page_break()
doc.add_heading('Layer responsibilities and maturity', level=1)
layer_rows=[
('Experience','Provide clear case navigation, portfolio overview, workflow guidance and reviewable outputs.','Current case pages, workflow optimizer, reports and Excel export; cross-case views and trace visualization are planned.'),
('Workflow orchestration','Move work through intake, analysis, decision, proposal, verification and packaging; surface next action and due date.','Core workflow guidance and e-GP checklist exist; assignment, approvals, reminders and escalations are future work.'),
('Analysis and decision','Turn source material into structured requirements, evidence links, readiness, gaps, risks and decision options.','Extraction, compliance, CV matching, scorecard, risk/gap, clarification and verification services exist; scenario comparison is planned.'),
('Case domain','Keep the tender opportunity as the shared object tying sources, criteria, evidence, people, decisions and outputs together.','SQLAlchemy models cover core case and diagnostic entities. Versioned baselines and formal decision records should be added deliberately.'),
('Trust and governance','Make findings reproducible and reviewable; preserve source lineage, uncertainty, human acceptance and output readiness.','Anchors, audit events, human-review language, readiness states and prohibited-phrase checks exist; source page-level citations, rule versioning and access policies need expansion.'),
('Data and knowledge','Store source files, extracted text, case records and reusable donor/sector/proposal intelligence.','Local SQLite/files and IP capture exist; indexing/search, retention controls, encryption and immutable snapshots are future work.'),
('Platform and operations','Provide reliable local operation first, then add service controls only when justified by product scope.','Local Streamlit/Python/SQLite MVP exists. Authentication, multi-tenant SaaS, billing and external AI are intentionally not implemented.'),
]
lt=doc.add_table(rows=1,cols=3); lt.alignment=WD_TABLE_ALIGNMENT.CENTER; lt.autofit=False
lwidth=[1.55,3.25,5.0]
for c,w,txt in zip(lt.rows[0].cells,lwidth,['Layer','Responsibility','Current state / direction']): c.width=Inches(w); set_cell_text(c,txt,True,(255,255,255),8.2); shade(c,'145270'); set_cell_margins(c)
set_repeat_table_header(lt.rows[0])
for i,row in enumerate(layer_rows):
    row_obj=lt.add_row(); prevent_row_split(row_obj); cs=row_obj.cells
    for c,w,txt in zip(cs,lwidth,row): c.width=Inches(w); set_cell_text(c,txt,size=8); set_cell_margins(c)
    for c in cs: shade(c,'F2F6F8' if i%2==0 else 'FFFFFF')

# Next page for development framework
doc.add_page_break()
doc.add_heading('Project development framework', level=1)
doc.add_paragraph('Use the Tenderly pattern as a sequence of product questions: what must the operator see, what should be testable before commitment, what needs monitoring, and what foundation makes it dependable? For TenderOS, translate that sequence into a governed tender lifecycle and build in increments.')
phases=[
('0. Product boundary and evidence baseline','Confirm users, paid diagnostic service, case lifecycle, local-first deployment, supported tender contexts and exclusions. Capture existing screens, data entities, output examples and failure cases.','Approved concept, source-of-truth inventory, capability register with implemented/partial/proposed tags.','MVP scope already excludes public SaaS, billing, multi-tenant auth and connected AI.'),
('1. Case console and source spine','Make case workspace the stable entry point. Improve source register, amendment/version handling, requirement anchors, global search and activity timeline.','Every finding points to a case, source version and requirement; users can find the latest authoritative input.','Preserve original files; flag OCR uncertainty and require human review.'),
('2. Simulation and decision workbench','Add saved scenarios for bid/no-bid alternatives; make assumptions editable; calculate score/risk changes; compare variants to a locked baseline.','A reviewer can explain why recommendation changed and reproduce each scenario.','No scenario may overwrite evidence or imply certainty; show unknowns and sensitivity.'),
('3. Traceable analysis and reusable modules','Build end-to-end requirement -> evidence -> gap/risk -> proposal section trace. Govern donor/context packs, methodology modules and reusable IP.','Coverage view; source-linked findings; reviewed reusable modules with owner, version and applicability.','Prevent unsupported claims; retain source context and avoid cross-client data leakage.'),
('4. Monitoring and controlled actions','Introduce deadline and addendum tracking, missing-input alerts, stale-data signals and assigned tasks. Add in-app routing before external integrations.','Exceptions are visible, owned, due-dated and auditable; notification delivery can be reconciled.','External messages, client-facing changes and bid decisions remain human-approved.'),
('5. Verification, packaging and operations','Expand reproducibility, role-based review, output versioning, backup/restore and operational checks. Add hosted/auth/API capability only if deployment model changes.','Repeatable client-ready package with readiness reason, audit trail, source trace and recovery procedure.','Keep local-only as default until security, privacy, support and operating model are specified.'),
]
pt=doc.add_table(rows=1,cols=4); pt.alignment=WD_TABLE_ALIGNMENT.CENTER; pt.autofit=False
pwidth=[1.5,3.45,2.65,2.2]
for c,w,txt in zip(pt.rows[0].cells,pwidth,['Phase','Work','Exit evidence','Design guardrail']): c.width=Inches(w); set_cell_text(c,txt,True,(255,255,255),8.1); shade(c,'145270'); set_cell_margins(c)
set_repeat_table_header(pt.rows[0])
for i,row in enumerate(phases):
    row_obj=pt.add_row(); prevent_row_split(row_obj); cs=row_obj.cells
    for c,w,txt in zip(cs,pwidth,row): c.width=Inches(w); set_cell_text(c,txt,size=7.7); set_cell_margins(c)
    for c in cs: shade(c,'F2F6F8' if i%2==0 else 'FFFFFF')

doc.add_heading('Core product rules', level=2)
for item in [
'Evidence before assertion: each compliance conclusion, risk and CV claim should resolve to supplied evidence or be marked Gap/Pending Input.',
'Scenarios are reversible: save assumptions and variants, show the baseline, and explain score changes instead of mutating the case silently.',
'Human approval at consequential boundaries: bid/no-bid, client-facing language, evidence acceptance and external messages require a named reviewer.',
'Reproducibility by default: preserve source versions, normalized requirements, rule versions, reviewer actions, readiness result and output package identity.',
'Local-first remains a product constraint until a separate decision establishes access control, hosting, retention, support and privacy requirements.'
]: doc.add_paragraph(item,style='List Bullet')

doc.add_heading('Measures of progress', level=2)
doc.add_paragraph('Track operational quality rather than win-rate claims: percentage of mandatory requirements linked to a source and evidence state; unresolved mandatory gaps at decision time; time from source intake to reviewed recommendation; extraction correction rate; stale-input/addendum acknowledgement time; verification failures before export; and the proportion of client-ready findings with a complete reviewer and source trail.')

p=doc.add_paragraph()
p.paragraph_format.space_before=Pt(5)
p.paragraph_format.space_after=Pt(0)
r=p.add_run('Sources and scope: Tenderly official documentation (https://docs.tenderly.co/introduction) and product overview (https://tenderly.co/pricing), checked 19 October 2026. TenderOS status is based on tenderproof_os/README.md, QUICKSTART.md, UI modules, models.py and services/ in this workspace; it describes this repository snapshot, not an independently deployed system.')
r.font.size=Pt(7.3); r.font.color.rgb=RGBColor(95,110,122)

# footer
for s in doc.sections:
    fp=s.footer.paragraphs[0]; fp.alignment=WD_ALIGN_PARAGRAPH.RIGHT
    fr=fp.add_run('Tenderly to TenderOS concept map  |  19 Oct 2026'); fr.font.name='Aptos'; fr.font.size=Pt(7); fr.font.color.rgb=RGBColor(110,125,137)

doc.core_properties.title='Tenderly to TenderOS Concept Map'
doc.core_properties.subject='Feature mapping, TenderOS architecture layers, and project development framework'
doc.core_properties.author='OpenAI'
doc.save(OUT)
print(OUT.resolve())
