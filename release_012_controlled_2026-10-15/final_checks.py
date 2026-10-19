"""Resumed verification procedures. No application code edits or gate approvals."""

import argparse
import hashlib
import json
import os
from pathlib import Path
import subprocess
import sys

sys.path.insert(0, str(Path(__file__).resolve().parent))
from control import APP, AUDIT, EVIDENCE, ROOT, digest, preserve, save
from verify import PYTHON, command, environment


def setup(label):
    os.environ.pop("TENDERPROOF_DB_PATH", None)
    os.environ.update(environment(label))
    sys.path.insert(0, str(APP))
    preserve()


def workflow():
    setup("final_workflow")
    from docx import Document
    from openpyxl import load_workbook
    from tenderproof import models as m
    from tenderproof.config import DB_PATH
    from tenderproof.db import get_session, init_db
    from tenderproof.services.approval_service import approve_case
    from tenderproof.services.compliance_service import update_compliance_row
    from tenderproof.services.decision_service import release_context
    from tenderproof.services.export_service import REQUIRED_SHEETS, export_case_to_excel
    from tenderproof.services.report_service import generate_docx_report, generate_markdown_report
    from tenderproof.services.verification_service import run_verification
    sys.path.insert(0, str(APP / "tests"))
    from test_release_012_workflow import reviewed_assessment

    if DB_PATH.exists():
        raise RuntimeError("Refusing to overwrite an existing verification runtime.")
    init_db()
    folder = ROOT / "runtime/final_workflow/supporting_evidence"
    folder.mkdir(parents=True)
    with get_session() as session:
        case, source, requirement, row, evidence, proof = reviewed_assessment(session, folder)
        case.tender_title = "SYNTHETIC QA - approved license diagnostic"
        case.country = "Bangladesh"
        case.sector = "Synthetic qualification review"
        session.add(m.CVMatch(expert_match_id="QA-CV", case_id=case.case_id, expert_role="Diagnostic reviewer", candidate_name="Synthetic expert", cv_evidence="Synthetic CV supplied for workflow testing only.", match_score_1_to_5=3))
        session.add(m.ProposalSkeleton(case_id=case.case_id, section_order=1, section_title="Understanding the assignment", purpose="Evidence-grounded synthetic scope", required_inputs="Reviewed tender source"))
        session.add(m.MethodologyOutline(case_id=case.case_id, module_order=1, module_title="Evidence review", module_description="Reconcile tender clauses with provided documents.", data_collection_conditions="Use only synthetic supplied evidence.", analysis_approach="Deterministic review"))
        session.add(m.RiskRegister(risk_id="QA-RISK", case_id=case.case_id, risk_category="Evidence", risk_description="Natural evidence expiry", severity="low", likelihood="low", mitigation="Recheck validity before each publication", owner="Synthetic reviewer", status="open"))
        initial = run_verification(session, case.case_id)
        if initial.blockers or any(r.rule_code != "APP-001" for r in initial.warnings):
            raise RuntimeError("STOP: valid synthetic inputs did not achieve technical verification.")
        approval = approve_case(session, case.case_id, reviewer="SYNTHETIC QA approver - not a real bid authorization", rationale="Explicit simulated review of this exact synthetic evidence snapshot for software testing.")
        case_id, approval_id = case.case_id, approval.approval_id
    with get_session() as session:
        context = release_context(session, case_id)
        if not context["approved"] or context["eligibility"] != "PASS":
            raise RuntimeError("STOP: approval did not survive database persistence.")
        markdown = generate_markdown_report(session, case_id)
        docx = generate_docx_report(session, case_id)
        xlsx = export_case_to_excel(session, case_id)
    with get_session() as session:
        result = run_verification(session, case_id)
        if result.overall_status != "passed" or result.readiness_label != "EXPORT_READY":
            raise RuntimeError("STOP: persisted approved deliverables did not verify.")
    md_text = markdown.read_text(encoding="utf-8")
    document = Document(docx)
    doc_text = "\n".join(p.text for p in document.paragraphs)
    workbook = load_workbook(xlsx, data_only=False)
    summary = dict(zip([c.value for c in workbook["Case Summary"][1]], [c.value for c in workbook["Case Summary"][2]]))
    complete_sections = all(f"## {i}. " in md_text for i in range(1, 16))
    good = complete_sections and "Prepared for internal/client review." in md_text and "Generated:" in md_text and "CLIENT_READY" in doc_text and approval_id in doc_text and workbook.sheetnames == REQUIRED_SHEETS and summary["readiness_label"] == "EXPORT_READY" and summary["approval_id"] == approval_id
    workbook.close()
    if not good:
        raise RuntimeError("STOP: actual MD/DOCX/XLSX artifact read-back failed.")

    negative_folder = folder / "negative"
    negative_folder.mkdir()
    with get_session() as session:
        negative, _, _, comp, *_ = reviewed_assessment(session, negative_folder)
        negative.tender_title = "SYNTHETIC QA - blocked diagnostic"
        run_verification(session, negative.case_id)
        approve_case(session, negative.case_id, reviewer="Synthetic QA", rationale="Simulated valid starting point.")
        old = generate_markdown_report(session, negative.case_id)
        old_bytes = old.read_bytes()
        update_compliance_row(session, comp.compliance_id, status="Gap", available_evidence="")
        blocked = release_context(session, negative.case_id)
        if blocked["approved"] or blocked["readiness_label"] != "BLOCKED":
            raise RuntimeError("STOP: evidence removal failed to block release.")
        draft = generate_markdown_report(session, negative.case_id)
        if "Readiness label: **BLOCKED**" not in draft.read_text() or old.read_bytes() != old_bytes:
            raise RuntimeError("STOP: blocked diagnostic or historical artifact preservation failed.")
        negative_id = negative.case_id
    save("positive_workflow.json", {"synthetic_only": True, "real_human_or_bid_approval_claimed": False, "case_id": case_id, "blocked_case_id": negative_id, "approval_id": approval_id, "persisted_separate_sessions": True, "readiness": "EXPORT_READY", "all_14_sections_plus_egp": complete_sections, "workbook_sheets": REQUIRED_SHEETS, "artifacts": [{"path": str(p.relative_to(ROOT)), "sha256": digest(p), "bytes": p.stat().st_size} for p in (markdown, docx, xlsx)], "blocked_report": str(draft.relative_to(ROOT)), "historical_report_unchanged": True, "technical_observation_passed": good, "human_merge_approval": "pending_approval"})
    preserve()
    print(json.dumps({"case": case_id, "readiness": "EXPORT_READY", "MD_DOCX_XLSX_read_back": True, "separate_sessions": True, "negative_release_blocked": True}))


def ui():
    output = EVIDENCE / "page_smoke.json"
    record = command("page_smoke", [str(PYTHON), "-B", str(AUDIT / "ui_smoke.py"), str(output)], env=environment("final_page_smoke"))
    results = json.loads(output.read_text())
    if record["exit_code"] or len(results) != 16 or any(r["exceptions"] for r in results):
        raise RuntimeError("STOP: 16-page Streamlit verification failed.")
    print("Fresh AppTest: all 16 existing pages rendered without exceptions.")


def launchers():
    env = environment("launcher_checks")
    env["TENDERPROOF_LAUNCHER_VERIFY"] = "1"
    env["PATH"] = str(APP / ".venv/Scripts") + os.pathsep + env["PATH"]
    results = []
    plans = [
        ("pip_check", [str(PYTHON), "-m", "pip", "check"]),
        ("cli_check", [str(PYTHON), "-B", "run_app.py", "--check"]),
        ("cli_init", [str(PYTHON), "-B", "run_app.py", "--init-db"]),
        ("cli_seed", [str(PYTHON), "-B", "run_app.py", "--seed-demo"]),
        ("primary_launcher", ["powershell.exe", "-NoProfile", "-NonInteractive", "-Command", "& '.\\deploy&run.bat'\nexit $LASTEXITCODE"]),
        ("safe_launcher", ["powershell.exe", "-NoProfile", "-NonInteractive", "-Command", "& '.\\deploy_and_run.bat'\nexit $LASTEXITCODE"]),
    ]
    for label, args in plans:
        record = command(label, args, env=env)
        results.append(record)
        if record["exit_code"]:
            save("launcher_summary.json", results)
            raise RuntimeError(f"STOP: launcher/check failed: {label}")
    save("launcher_summary.json", {"commands": results, "batch_mode": "explicit health-only verification mode; no deployment", "clean_install_log": "evidence/clean_install.log", "full_default_interactive_double_click_executed": False})
    print("CLI flags, pip check and both exact batch filenames passed isolated health-mode execution.")


def transactions():
    setup("transactions")
    from sqlalchemy.exc import IntegrityError
    from tenderproof.db import get_session, init_db
    from tenderproof.models import Case
    from tenderproof.services.case_service import create_case

    path = ROOT / "runtime/transactions/storage/transactions.db"
    if path.exists():
        raise RuntimeError("Refusing to overwrite an existing transaction-verification DB.")
    init_db(path)
    with get_session(path) as session:
        create_case(session, case_id="TX-KEEP", client_name="Synthetic rollback baseline")
    conflict_observed = False
    try:
        with get_session(path) as session:
            create_case(session, case_id="TX-PARTIAL", client_name="Must roll back")
            create_case(session, case_id="TX-KEEP", client_name="Duplicate synthetic ID")
    except IntegrityError:
        conflict_observed = True
    with get_session(path) as session:
        partial_absent = session.query(Case).filter_by(case_id="TX-PARTIAL").count() == 0
        baseline_intact = session.query(Case).filter_by(case_id="TX-KEEP").one().client_name == "Synthetic rollback baseline"
        create_case(session, case_id="TX-AFTER", client_name="Healthy session after conflict")
    with get_session(path) as session:
        subsequent_commit = session.query(Case).filter_by(case_id="TX-AFTER").count() == 1
    result = {"synthetic_only": True, "unique_conflict_observed": conflict_observed, "partial_transaction_rolled_back": partial_absent, "baseline_record_preserved": baseline_intact, "subsequent_session_commit_passed": subsequent_commit}
    save("transactions.json", result)
    if not all(result.values()):
        raise RuntimeError("STOP: transaction rollback verification failed.")
    preserve()
    print("Synthetic unique conflict rolled back partial writes; baseline and subsequent session remained intact.")


if __name__ == "__main__":
    p = argparse.ArgumentParser()
    p.add_argument("phase", choices=["workflow", "ui", "launchers", "transactions"])
    a = p.parse_args()
    {"workflow": workflow, "ui": ui, "launchers": launchers, "transactions": transactions}[a.phase]()
