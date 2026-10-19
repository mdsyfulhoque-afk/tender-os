"""Discriminator tests start from valid evidence, not merely missing approval."""

from datetime import UTC, datetime, timedelta
import json
from pathlib import Path
from types import SimpleNamespace

from docx import Document
from openpyxl import load_workbook
from PIL import Image
import pytest

from tenderproof import models as m
from tenderproof.constants import REQUIRED_BID_CATEGORIES
from tenderproof.db import make_session
from tenderproof.services import decision_service, export_service, report_service, source_service
from tenderproof.services.approval_service import approve_case, record_evidence_review, review_source
from tenderproof.services.bid_score_service import save_scores_and_recommendation
from tenderproof.services.case_service import create_case
from tenderproof.services.compliance_service import update_compliance_row
from tenderproof.services.decision_service import mandatory_state, policy_assessment, release_context, snapshot_hash
from tenderproof.services.extraction_service import update_requirement
from tenderproof.services.verification_service import _scan_outputs_for_prohibited_phrases, run_verification
from tenderproof.services.workflow_service import build_workflow_steps

from test_release_012_workflow import approved_assessment, reviewed_assessment


def test_reviewed_not_applicable_is_explicit_and_can_be_released(session, tmp_path):
    case, source, _, row, *_ = reviewed_assessment(session, tmp_path)
    update_compliance_row(session, row.compliance_id, status="Not Applicable", available_evidence="")
    assert mandatory_state(session, case.case_id) == "UNKNOWN"
    record_evidence_review(session, case.case_id, row.compliance_id, reviewer="Synthetic applicability reviewer", notes="Synthetic clause explicitly excludes this bidder category; checked original source.", legal_entity=case.organization, source_id=source.source_id, review_status="not_applicable")
    assert mandatory_state(session, case.case_id, verified=True) == "NOT_APPLICABLE"
    result = run_verification(session, case.case_id)
    assert not result.blockers
    approve_case(session, case.case_id, reviewer="Synthetic human", rationale="Reviewed documented applicability exclusion.")
    assert release_context(session, case.case_id)["approved"]
    assert "Eligibility: NOT_APPLICABLE" in report_service.generate_markdown_report(session, case.case_id).read_text()


@pytest.mark.parametrize("state,expected", [("failed", "FAIL"), ("pending", "UNKNOWN"), ("verified", "PASS")])
def test_explicit_evidence_outcomes_have_distinct_semantics(session, tmp_path, state, expected):
    case, _, _, row, _, proof = reviewed_assessment(session, tmp_path)
    record_evidence_review(session, case.case_id, row.compliance_id, reviewer="Synthetic reviewer", notes="Explicit controlled evidence outcome.", legal_entity=case.organization, file_path=proof, validity_reason="Synthetic timeless supporting fact.", review_status=state)
    assert mandatory_state(session, case.case_id, verified=True) == expected
    scores = {c: {"score": 5} for c in REQUIRED_BID_CATEGORIES}
    rec = save_scores_and_recommendation(session, case.case_id, scores)
    if expected == "FAIL":
        assert rec.recommendation == "Reject / No-bid"


@pytest.mark.parametrize("mutation,rule", [("placeholder", "EVID-001"), ("expired", "EVID-002"), ("wrong_entity", "EVID-002"), ("source_partial", "SRC-003"), ("coverage", "SRC-003"), ("requirement", "ELIG-001"), ("duplicate_compliance", "ELIG-001")])
def test_each_mandatory_barrier_rejects_previously_approved_case(session, tmp_path, mutation, rule):
    case, source, req, row, review, *_ = approved_assessment(session, tmp_path)
    assert release_context(session, case.case_id)["approved"]
    if mutation == "placeholder":
        row.available_evidence = "Pending Input"
    elif mutation == "expired":
        review.valid_until = "2020-01-01"
    elif mutation == "wrong_entity":
        review.legal_entity = "Another bidder"
    elif mutation == "source_partial":
        source.extraction_status = "partial"
    elif mutation == "coverage":
        source.extraction_coverage_json = json.dumps({"complete": False, "processed_pages": 1, "total_pages": 2})
    elif mutation == "requirement":
        req.needs_review = True
    else:
        session.add(m.ComplianceLedger(compliance_id="DUPLICATE-SYNTHETIC", case_id=case.case_id, requirement_id=req.requirement_id, status="Met", available_evidence="Duplicated interpretation"))
    session.flush()
    policy = policy_assessment(session, case.case_id)
    assert any(code == rule for code, _ in policy["blockers"])
    assert not release_context(session, case.case_id)["approved"]
    with pytest.raises(ValueError):
        approve_case(session, case.case_id, reviewer="Synthetic human", rationale="Attempt to override technical blocker.")


@pytest.mark.parametrize("entity", ["ip", "cv", "question", "method", "case", "score", "risk", "gap", "skeleton"])
def test_report_input_changes_cannot_reuse_old_approval(session, tmp_path, entity):
    case, *_ = approved_assessment(session, tmp_path)
    before = snapshot_hash(session, case.case_id)
    if entity == "ip":
        session.add(m.IPCapture(ip_id="IP-CHANGE", case_id=case.case_id, pattern_title="New pattern", pattern_type="Synthetic", anonymized=True))
    elif entity == "cv":
        session.add(m.CVMatch(expert_match_id="CV-CHANGE", case_id=case.case_id, candidate_name="Synthetic person", cv_evidence="Provided history", match_score_1_to_5=3))
    elif entity == "question":
        session.add(m.ClarificationQuestion(question_id="Q-CHANGE", case_id=case.case_id, question="New clarification?"))
    elif entity == "method":
        session.add(m.MethodologyOutline(case_id=case.case_id, module_order=1, module_title="Changed method"))
    elif entity == "case":
        case.submission_deadline = "2099-01-01"
    elif entity == "score":
        session.query(m.BidScorecard).filter_by(case_id=case.case_id).first().score = 1
    elif entity == "risk":
        session.add(m.RiskRegister(risk_id="R-CHANGE", case_id=case.case_id, risk_description="New risk"))
    elif entity == "gap":
        session.add(m.GapItem(gap_id="G-CHANGE", case_id=case.case_id, gap_description="New input"))
    else:
        session.add(m.ProposalSkeleton(case_id=case.case_id, section_order=1, section_title="Changed structure"))
    session.flush()
    assert snapshot_hash(session, case.case_id) != before
    assert not release_context(session, case.case_id)["approved"]


def test_amendment_review_reassessment_and_reapproval_can_recover(session, tmp_path):
    case, source, req, row, *_ = approved_assessment(session, tmp_path)
    old_hash = snapshot_hash(session, case.case_id)
    newer = source_service.register_manual_text(session, case.case_id, "Bidder must hold a current license. Amendment: review the new closing time.", filename=source.filename)
    assert not release_context(session, case.case_id)["approved"]
    update_requirement(session, req.requirement_id, source_id=newer.source_id, needs_review=False, notes="Reviewed unchanged license clause in amended source.")
    review_source(session, case.case_id, newer.source_id, reviewer="Synthetic amendment reviewer", notes="Checked all amendment clauses; licensing requirement unchanged and closing date reconciled.", requirements_complete=True)
    run_verification(session, case.case_id)
    approve_case(session, case.case_id, reviewer="Synthetic human", rationale="Reviewed amended source and unchanged attached license evidence.")
    assert release_context(session, case.case_id)["approved"]
    assert snapshot_hash(session, case.case_id) != old_hash
    assert session.query(m.DecisionApproval).count() == 2


@pytest.mark.parametrize("format_name", ["markdown", "docx", "xlsx"])
def test_revocation_during_generation_quarantines_without_publication(session, tmp_path, monkeypatch, format_name):
    case, *_, approval = approved_assessment(session, tmp_path)
    case_id, approval_id = case.case_id, approval.approval_id
    session.commit()
    original_hash = snapshot_hash(session, case_id)
    factory = make_session(session.bind.url.database)

    def revoke_from_separate_connection():
        with factory() as other:
            other.query(m.DecisionApproval).filter_by(approval_id=approval_id).one().revoked_at = datetime.now(UTC)
            other.commit()

    if format_name == "xlsx":
        original = export_service.validate_publication

        def hooked(*args):
            revoke_from_separate_connection()
            return original(*args)

        monkeypatch.setattr(export_service, "validate_publication", hooked)
        action = export_service.export_case_to_excel
        directory = export_service.EXPORTS_DIR
    else:
        original = report_service.build_report_markdown

        def hooked(*args):
            content = original(*args)
            revoke_from_separate_connection()
            return content

        monkeypatch.setattr(report_service, "build_report_markdown", hooked)
        action = report_service.generate_markdown_report if format_name == "markdown" else report_service.generate_docx_report
        directory = report_service.REPORTS_DIR
    with pytest.raises(ValueError, match="Approval"):
        action(session, case_id)
    assert snapshot_hash(session, case_id) == original_hash, "Revocation itself is not a content mutation"
    assert session.query(m.OutputPackage).filter_by(case_id=case_id).count() == 0
    assert list((directory / "quarantine").glob(f"{case_id}*"))


@pytest.mark.parametrize("where", ["table", "header", "footer", "unreadable", "missing"])
def test_safe001_inspects_nonparagraph_content_and_fails_closed(tmp_path, where):
    path = tmp_path / "content.docx"
    doc = Document()
    if where == "table":
        doc.add_table(rows=1, cols=1).cell(0, 0).text = "GUARANTEED WIN"
    elif where in {"header", "footer"}:
        getattr(doc.sections[0], where).paragraphs[0].text = "guaranteed success"
    doc.save(path)
    if where == "unreadable":
        path.write_bytes(b"broken document")
    elif where == "missing":
        path.unlink()
    assert _scan_outputs_for_prohibited_phrases([m.OutputPackage(output_type="docx_report", file_path=str(path))])


@pytest.mark.parametrize("value", ["=1+1", "+SUM(1,1)", "-1+2", "@SUM(1,1)"])
def test_formula_shaped_strings_remain_literal_and_numbers_remain_numeric(session, tmp_path, value):
    case, *_ = reviewed_assessment(session, tmp_path)
    case.client_name = value
    book = load_workbook(export_service.export_case_to_excel(session, case.case_id), data_only=False)
    sheet = book["Case Summary"]
    col = next(c.column for c in sheet[1] if c.value == "client_name")
    assert sheet.cell(2, col).value == value
    assert sheet.cell(2, col).data_type == "s"
    score = book["Bid No-Bid Scorecard"]
    index = next(c.column for c in score[1] if c.value == "score")
    assert score.cell(2, index).data_type == "n"


def test_source_coverage_persists_docx_tables_and_mixed_pdf(tmp_path, monkeypatch):
    path = tmp_path / "table.docx"
    doc = Document()
    doc.add_table(rows=1, cols=1).cell(0, 0).text = "Bidder must provide table-only license evidence."
    doc.save(path)
    coverage = {}
    text, status, _ = source_service.extract_text(path, coverage=coverage)
    assert status == "extracted" and "table-only" in text and coverage["tables"] == 1

    import pypdfium2
    import rapidocr

    class PDF:
        def __len__(self):
            return 2

        def __getitem__(self, index):
            return SimpleNamespace(render=lambda scale: SimpleNamespace(to_pil=lambda: Image.new("RGB", (80, 60))))

    calls = []
    monkeypatch.setattr(source_service, "PdfReader", lambda path: SimpleNamespace(pages=[SimpleNamespace(extract_text=lambda: "Native page with mandatory qualification text."), SimpleNamespace(extract_text=lambda: "")]))
    monkeypatch.setattr(pypdfium2, "PdfDocument", lambda path: PDF())
    monkeypatch.setattr(rapidocr, "RapidOCR", lambda: lambda image: calls.append(True) or SimpleNamespace(txts=["Scanned page with additional obligations."]))
    coverage = {}
    text, status, _ = source_service.extract_text(tmp_path / "mixed.pdf", coverage=coverage)
    assert status == "extracted" and coverage["complete"]
    assert "Native page" in text and "Scanned page" in text and len(calls) == 1
    capped = {}
    _, note = source_service._ocr_pdf(tmp_path / "mixed.pdf", max_pages=1, coverage=capped)
    assert not capped["complete"] and capped["processed_pages"] == 1 and "limited to first" in note


def test_workflow_does_not_complete_tampered_or_draft_outputs(session, tmp_path):
    case, *_ = approved_assessment(session, tmp_path)
    report = report_service.generate_markdown_report(session, case.case_id)
    export_service.export_case_to_excel(session, case.case_id)
    steps = {s["step"]: s["complete"] for s in build_workflow_steps(session, case)}
    assert steps["Generate diagnostic report"] and steps["Export workbook"]
    report.write_text("Altered report", encoding="utf-8")
    steps = {s["step"]: s["complete"] for s in build_workflow_steps(session, case)}
    assert not steps["Generate diagnostic report"]


def test_create_cannot_assign_release_clearance(session):
    case = create_case(session, client_name="Synthetic bypass", readiness_label="EXPORT_READY", status="packaged")
    assert case.readiness_label == "DRAFT_ONLY" and case.status == "draft"


def test_ip_csv_does_not_publish_formula_commands(session):
    import csv
    from tenderproof.services.ip_capture_service import add_ip_capture, export_ip_csv

    case = create_case(session, client_name="Synthetic IP")
    add_ip_capture(session, case.case_id, pattern_title="=1+1", pattern_description="Ordinary text")
    with export_ip_csv(session, case.case_id).open(encoding="utf-8", newline="") as stream:
        value = next(csv.DictReader(stream))
    assert value["pattern_title"] == "'=1+1" and value["pattern_description"] == "Ordinary text"
