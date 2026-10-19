from pathlib import Path
from uuid import uuid4
import json

from docx import Document
from sqlalchemy.orm import Session

from ..constants import PROHIBITED_REPORT_PHRASES, REQUIRED_BID_CATEGORIES
from ..models import (
    BidRecommendation,
    BidScorecard,
    Case,
    ComplianceLedger,
    CVMatch,
    GapItem,
    NormalizedRequirement,
    OutputPackage,
    RiskRegister,
    Source,
    VerificationResult,
    VerificationRun,
)
from ..schemas import RuleResult, VerificationSummary
from .audit_service import log_event
from .decision_service import current_approval, evidence_review_problem, latest_evidence_review, mandatory_state, output_is_current, policy_assessment, snapshot_data, snapshot_hash


def run_verification(session: Session, case_id: str) -> VerificationSummary:
    session.flush()
    results: list[RuleResult] = []

    case = session.query(Case).filter(Case.case_id == case_id).first()
    _add(results, "CASE-001", "Case exists", bool(case), "blocker", f"Case {case_id} exists." if case else "Case does not exist.")
    if not case:
        return _store_and_return(session, case_id, None, results)

    sources = session.query(Source).filter(Source.case_id == case_id).all()
    reqs = session.query(NormalizedRequirement).filter(NormalizedRequirement.case_id == case_id).all()
    comp = session.query(ComplianceLedger).filter(ComplianceLedger.case_id == case_id).all()
    gaps = session.query(GapItem).filter(GapItem.case_id == case_id).all()
    cvs = session.query(CVMatch).filter(CVMatch.case_id == case_id).all()
    risks = session.query(RiskRegister).filter(RiskRegister.case_id == case_id).all()
    score_rows = session.query(BidScorecard).filter(BidScorecard.case_id == case_id).all()
    rec = session.query(BidRecommendation).filter(BidRecommendation.case_id == case_id).order_by(BidRecommendation.updated_at.desc()).first()
    outputs = session.query(OutputPackage).filter(OutputPackage.case_id == case_id).all()

    _add(results, "SRC-001", "Case has ToR/source input", bool(sources), "blocker", "At least one source exists." if sources else "No ToR/source file or manual ToR text is registered.")
    _add(results, "SRC-002", "Sources have extraction status", all(bool(s.extraction_status) for s in sources), "warning", "All sources have extraction status." if sources else "No sources to verify.")
    _add(results, "REQ-001", "Case has normalized requirements", bool(reqs), "blocker", "Requirements exist." if reqs else "No normalized requirements exist.")
    _add(results, "REQ-002", "Requirements have keys and anchors", all(r.match_key and r.canonical_anchor for r in reqs), "blocker", "All requirements have match_key and canonical_anchor.")
    mandatory_ids = {r.requirement_id for r in reqs if r.mandatory}
    comp_req_ids = {c.requirement_id for c in comp}
    _add(results, "REQ-003", "Mandatory requirements have compliance rows", mandatory_ids.issubset(comp_req_ids), "blocker", "Mandatory requirements have compliance rows.")

    gap_comp = [c for c in comp if c.status in {"Gap", "Pending Input"}]
    gap_comp_ids = {g.related_compliance_id for g in gaps if g.related_compliance_id}
    _add(results, "COMP-001", "Gap and Pending Input compliance rows create gaps", all(c.compliance_id in gap_comp_ids for c in gap_comp), "warning", "Compliance gaps are represented as GapItems.")
    comp_by_req = {c.requirement_id: c for c in comp}
    no_evidence_failures = []
    for req in reqs:
        row = comp_by_req.get(req.requirement_id)
        if row and row.status == "Not Applicable":
            review = latest_evidence_review(session, row)
            if review and review.review_status == "not_applicable" and not evidence_review_problem(session, case, row, review):
                continue
        if req.mandatory and row and not (row.available_evidence or "").strip() and row.status not in {"Gap", "Pending Input"}:
            no_evidence_failures.append(req.requirement_id)
    _add(results, "COMP-002", "Mandatory requirements without evidence are gaps or pending", not no_evidence_failures, "blocker", "Mandatory missing evidence is correctly marked.")

    bad_cv = [cv.expert_match_id for cv in cvs if not (cv.cv_evidence or "").strip() and (cv.match_score_1_to_5 or 0) >= 4]
    _add(results, "CV-001", "CV evidence is not invented", not bad_cv, "blocker", "No high CV scores without evidence." if not bad_cv else f"High CV score without evidence: {', '.join(bad_cv)}")

    high_req_ids = {r.requirement_id for r in reqs if r.risk_level in {"high", "critical"}}
    high_req_ids |= {c.requirement_id for c in comp if c.risk_level in {"high", "critical"}}
    high_gap_req_ids = {g.related_requirement_id for g in gaps if g.status != "closed" and (g.related_requirement_id in high_req_ids or "high" in (g.gap_type or "").lower() or "critical" in (g.gap_type or "").lower())}
    risk_req_ids = {r.source_requirement_id for r in risks if r.source_requirement_id}
    _add(results, "RISK-001", "High or critical gaps appear in risk register", high_gap_req_ids.issubset(risk_req_ids), "warning", "High/critical gaps are reflected in risk register.")

    categories = {row.category for row in score_rows}
    _add(results, "BID-001", "Bid scorecard has all required categories", set(REQUIRED_BID_CATEGORIES).issubset(categories), "blocker", "Bid scorecard has all required categories.")
    _add(results, "BID-002", "Bid recommendation exists", bool(rec), "blocker", "Bid recommendation exists." if rec else "Bid recommendation missing.")

    has_report = any(o.output_type in {"markdown_report", "docx_report"} and Path(o.file_path).exists() for o in outputs)
    has_excel = any(o.output_type == "excel_export" and Path(o.file_path).exists() for o in outputs)
    _add(results, "EXP-001", "Diagnostic report exists before export-ready", has_report or case.readiness_label not in {"CLIENT_READY", "EXPORT_READY"}, "blocker", "Diagnostic report exists." if has_report else "Diagnostic report missing.")
    _add(results, "EXP-002", "Excel export exists before export-ready", has_excel or case.readiness_label != "EXPORT_READY", "blocker", "Excel export exists." if has_excel else "Excel export missing.")

    fingerprint = snapshot_hash(session, case_id)
    relevant_outputs = [o for o in outputs if o.verification_status != "stale" and o.snapshot_hash in {None, fingerprint}]
    prohibited = _scan_outputs_for_prohibited_phrases(relevant_outputs)
    _add(results, "SAFE-001", "Report avoids prohibited claims", not prohibited, "blocker", "No prohibited guarantee claims found." if not prohibited else f"Prohibited phrases found: {', '.join(prohibited)}")

    policy = policy_assessment(session, case_id)
    _add(results, "ELIG-001", "Mandatory requirements reviewed and resolved", mandatory_state(session, case_id) in {"PASS", "NOT_APPLICABLE"}, "blocker", "Mandatory requirements need reviewed evidence or an attributable source-based exclusion.")
    for code, message in policy["blockers"]:
        if code != "ELIG-001":
            _add(results, code, "Release policy", False, "blocker", message)
    for code, message in policy["warnings"]:
        _add(results, code, "Pending human review", False, "warning", message)
    approval = current_approval(session, case_id, fingerprint)
    _add(results, "APP-001", "Current human approval", bool(approval), "warning", "Current snapshot has an approval." if approval else "A named human approval is required before client-ready release.")
    current_outputs = [o for o in outputs if output_is_current(o, fingerprint, approval)]
    has_report = any(o.output_type in {"markdown_report", "docx_report"} for o in current_outputs)
    has_excel = any(o.output_type == "excel_export" for o in current_outputs)
    readiness = _compute_readiness(results, has_report, has_excel, approved=bool(approval))
    _add(results, "LABEL-001", "Readiness label matches verification outcome", True, "info", f"Readiness label set to {readiness}.")
    case.readiness_label = readiness
    if readiness == "BLOCKED":
        case.status = "blocked"
    elif readiness == "PRELIMINARY":
        case.status = "preliminary"
    elif readiness in {"CLIENT_READY", "EXPORT_READY"}:
        case.status = "client_ready" if readiness == "CLIENT_READY" else "packaged"
    summary = _store_and_return(session, case_id, case, results, readiness, fingerprint)
    summary.quality_metrics["eligibility"] = policy["eligibility"]
    log_event(session, case_id, "verification.run", f"Verification {summary.overall_status}; readiness {summary.readiness_label}")
    return summary


def _add(results: list[RuleResult], code: str, name: str, passed: bool, severity: str, message: str) -> None:
    if passed:
        status = "passed"
        blocker = False
        sev = "info" if severity == "info" else severity
    else:
        status = "failed" if severity == "blocker" else "warning"
        blocker = severity == "blocker"
        sev = severity
    results.append(RuleResult(code, name, status, sev, message, blocker))


def _store_and_return(session: Session, case_id: str, case: Case | None, results: list[RuleResult], readiness: str | None = None, fingerprint: str | None = None) -> VerificationSummary:
    blockers = [r for r in results if r.blocker and r.status != "passed"]
    warnings = [r for r in results if r.status == "warning"]
    readiness_label = readiness or ("BLOCKED" if blockers else "PRELIMINARY" if warnings else "CLIENT_READY")
    overall = "failed" if blockers else "warning" if warnings else "passed"
    run_id = f"RUN-{uuid4().hex[:12].upper()}" if case else None
    if case:
        session.add(VerificationRun(run_id=run_id, case_id=case_id, snapshot_hash=fingerprint or snapshot_hash(session, case_id), snapshot_json=json.dumps(snapshot_data(session, case_id), sort_keys=True, default=str), technical_passed=all(r.status == "passed" for r in results if r.rule_code not in {"APP-001", "LABEL-001"}), overall_status=overall, readiness_label=readiness_label))
    for result in results if case else []:
        session.add(
            VerificationResult(
                verification_id=f"VER-{uuid4().hex[:10].upper()}",
                case_id=case_id,
                rule_code=result.rule_code,
                rule_name=result.rule_name,
                status=result.status,
                severity=result.severity,
                message=result.message,
                blocker=result.blocker,
                run_id=run_id,
            )
        )
    metrics = {
        "rules": len(results),
        "blockers": len(blockers),
        "warnings": len(warnings),
        "case_status": case.status if case else "missing",
        "run_id": run_id,
        "snapshot_hash": fingerprint,
    }
    session.flush()
    return VerificationSummary(overall, blockers, warnings, readiness_label, results, metrics)


def _compute_readiness(results: list[RuleResult], has_report: bool, has_excel: bool, approved: bool = False) -> str:
    blockers = [r for r in results if r.blocker and r.status != "passed"]
    warnings = [r for r in results if r.status == "warning"]
    if blockers:
        return "BLOCKED"
    if warnings or not approved:
        return "PRELIMINARY"
    if has_report and has_excel:
        return "EXPORT_READY"
    if has_report:
        return "CLIENT_READY"
    return "PRELIMINARY"


def _scan_outputs_for_prohibited_phrases(outputs: list[OutputPackage]) -> list[str]:
    found = set()
    for output in outputs:
        path = Path(output.file_path)
        if output.output_type not in {"markdown_report", "docx_report"}:
            continue
        if not path.is_file():
            found.add(f"Missing report: {path.name}")
            continue
        text = ""
        try:
            if path.suffix.lower() == ".md":
                text = path.read_text(encoding="utf-8")
            elif path.suffix.lower() == ".docx":
                doc = Document(str(path))
                chunks = [p.text for p in doc.paragraphs]
                chunks += [cell.text for table in doc.tables for row in table.rows for cell in row.cells]
                for section in doc.sections:
                    for part in (section.header, section.footer, section.first_page_header, section.first_page_footer, section.even_page_header, section.even_page_footer):
                        chunks += [p.text for p in part.paragraphs]
                        chunks += [cell.text for table in part.tables for row in table.rows for cell in row.cells]
                text = "\n".join(chunks)
            else:
                raise ValueError("Unsupported report format")
        except Exception:
            found.add(f"Unreadable report: {path.name}")
            continue
        lower = text.lower()
        for phrase in PROHIBITED_REPORT_PHRASES:
            if phrase in lower:
                found.add(phrase)
    return sorted(found)

