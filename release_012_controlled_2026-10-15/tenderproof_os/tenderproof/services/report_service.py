from datetime import UTC, datetime
from pathlib import Path
from uuid import uuid4

from docx import Document
from docx.shared import Pt
from sqlalchemy.orm import Session

from ..config import REPORTS_DIR, ensure_directories
from ..constants import PROHIBITED_REPORT_PHRASES
from ..models import (
    BidRecommendation,
    Case,
    ClarificationQuestion,
    ComplianceLedger,
    CVMatch,
    EGPChecklistItem,
    EvidenceReview,
    GapItem,
    MethodologyOutline,
    NormalizedRequirement,
    OutputPackage,
    ProposalSkeleton,
    RiskRegister,
    Source,
)
from .audit_service import log_event
from .decision_service import file_hash, quarantine_artifact, release_context, validate_publication


def generate_markdown_report(session: Session, case_id: str) -> Path:
    ensure_directories()
    case = session.query(Case).filter(Case.case_id == case_id).first()
    if not case:
        raise ValueError(f"Case not found: {case_id}")
    context = release_context(session, case_id, generating="report")
    content = build_report_markdown(session, case, context)
    path = REPORTS_DIR / f"{case_id}_TenderProof_Rapid_Bid_Diagnostic_{uuid4().hex[:10]}.pending.md"
    path.write_text(content, encoding="utf-8")
    output = _record_output(session, case_id, "markdown_report", path, context)
    log_event(session, case_id, "report.markdown", f"Generated Markdown report {path.name}")
    return Path(output.file_path)


def generate_docx_report(session: Session, case_id: str) -> Path:
    ensure_directories()
    case = session.query(Case).filter(Case.case_id == case_id).first()
    if not case:
        raise ValueError(f"Case not found: {case_id}")
    context = release_context(session, case_id, generating="report")
    md = build_report_markdown(session, case, context)
    path = REPORTS_DIR / f"{case_id}_TenderProof_Rapid_Bid_Diagnostic_{uuid4().hex[:10]}.pending.docx"
    doc = Document()
    normal = doc.styles["Normal"]
    normal.font.name = "Arial"
    normal.font.size = Pt(10.5)
    for line in md.splitlines():
        if line.startswith("# "):
            doc.add_heading(line[2:].strip(), level=0)
        elif line.startswith("## "):
            doc.add_heading(line[3:].strip(), level=1)
        elif line.startswith("### "):
            doc.add_heading(line[4:].strip(), level=2)
        elif line.startswith("- "):
            doc.add_paragraph(line[2:].strip(), style="List Bullet")
        elif line.strip() == "":
            doc.add_paragraph("")
        else:
            doc.add_paragraph(line)
    doc.save(path)
    output = _record_output(session, case_id, "docx_report", path, context)
    log_event(session, case_id, "report.docx", f"Generated DOCX report {path.name}")
    return Path(output.file_path)


def build_report_markdown(session: Session, case: Case, context=None) -> str:
    context = context or release_context(session, case.case_id, generating="report")
    sources = session.query(Source).filter(Source.case_id == case.case_id).all()
    reqs = session.query(NormalizedRequirement).filter(NormalizedRequirement.case_id == case.case_id).all()
    compliance = session.query(ComplianceLedger).filter(ComplianceLedger.case_id == case.case_id).all()
    cvs = session.query(CVMatch).filter(CVMatch.case_id == case.case_id).all()
    gaps = session.query(GapItem).filter(GapItem.case_id == case.case_id).all()
    risks = session.query(RiskRegister).filter(RiskRegister.case_id == case.case_id).all()
    questions = session.query(ClarificationQuestion).filter(ClarificationQuestion.case_id == case.case_id).all()
    skeleton = session.query(ProposalSkeleton).filter(ProposalSkeleton.case_id == case.case_id).order_by(ProposalSkeleton.section_order).all()
    methodology = session.query(MethodologyOutline).filter(MethodologyOutline.case_id == case.case_id).order_by(MethodologyOutline.module_order).all()
    egp_checks = session.query(EGPChecklistItem).filter(EGPChecklistItem.case_id == case.case_id).order_by(EGPChecklistItem.id).all()
    rec = session.query(BidRecommendation).filter(BidRecommendation.case_id == case.case_id).order_by(BidRecommendation.updated_at.desc()).first()

    met = sum(1 for row in compliance if row.status == "Met")
    gap_count = sum(1 for row in compliance if row.status in {"Gap", "Pending Input"})
    lines = [
        "# TenderProof Rapid Bid Diagnostic",
        "",
        f"Generated: {datetime.now(UTC).strftime('%Y-%m-%d %H:%M UTC')}",
        f"Readiness label: **{context['readiness_label']}**",
        f"Release status: {'APPROVED' if context['approved'] else 'BLOCKED' if context['blockers'] else 'DRAFT - human review or approval pending'}",
        f"Evidence snapshot: {context['snapshot_hash']}",
        f"Eligibility: {context['eligibility']}",
        f"Approval: {context['approval'].approval_id + ' by ' + context['approval'].approved_by if context['approved'] else 'Not approved'}",
        "",
        "Prepared for internal/client review.",
        "",
        "## 1. Client and tender summary",
        f"- Client: {case.client_name or 'Pending Input'}",
        f"- Organization: {case.organization or 'Pending Input'}",
        f"- Tender: {case.tender_title or 'Pending Input'}",
        f"- Donor/client: {case.donor_client or 'Pending Input'}",
        f"- Country / sector: {case.country or 'Pending Input'} / {case.sector or 'Pending Input'}",
        f"- Submission deadline: {case.submission_deadline or 'Pending Input'}",
        "",
        "## 2. Documents reviewed",
    ]
    lines += [f"- {src.filename or src.source_id} ({src.source_category or 'source'}, version: {src.version_number}; extraction: {src.extraction_status or 'Pending Input'}; human review: {src.reviewed_by or 'Pending Input'}; SHA-256: {src.content_sha256 or 'Legacy source - review pending'})" for src in sources] or ["- Pending Input: no source documents registered."]
    lines += [
        "",
        "## 3. Executive bid/no-bid recommendation",
    ]
    if rec:
        lines += [
            f"- Provisional opportunity recommendation: **{rec.recommendation}**",
            f"- Release decision: {'Approved' if context['approved'] else 'Blocked' if context['blockers'] else 'Review / approval pending'}",
            f"- Score: {rec.total_score}/{rec.max_score} ({rec.percentage}%)",
            f"- Rationale: {rec.rationale or 'Pending Input'}",
        ]
    else:
        lines.append("- Pending Input: bid/no-bid recommendation has not been generated.")
    lines += [
        "",
        "## 4. Opportunity summary",
        f"This diagnostic reviews the tender opportunity for compliance fit, expert/CV fit, proposal structuring, delivery risk, and bid/no-bid readiness. Missing evidence is marked as Gap or Pending Input.",
        "",
        "## 5. Compliance status summary",
        f"- Requirements normalized: {len(reqs)}",
        f"- Compliance rows: {len(compliance)}",
        f"- Met rows: {met}",
        f"- Gap or pending input rows: {gap_count}",
        "",
        "## 6. Key ToR requirements",
    ]
    lines += [f"- [{req.requirement_id}] {req.atomic_requirement_text} | Source: {req.source_id} | Page/section: {req.tor_section_page} | Anchor: {req.canonical_anchor} | Review: {'pending' if req.needs_review else 'reviewed'} | Risk: {req.risk_level}" for req in reqs] or ["- Pending Input: no normalized requirements."]
    lines += [
        "",
        "## 7. CV/expert match summary",
    ]
    lines += [f"- {cv.expert_role or 'Expert'} / {cv.candidate_name or 'Candidate'}: score {cv.match_score_1_to_5 or 0}/5. Evidence: {cv.cv_evidence or 'Pending Input'}. Gap: {cv.gap or 'None recorded'}" for cv in cvs] or ["- Pending Input: no CV match rows."]
    lines += [
        "",
        "## 8. Major gaps and pending inputs",
    ]
    lines += [f"- [{gap.gap_id}] {gap.gap_description} Action: {gap.required_action}" for gap in gaps] or ["- No open gaps recorded."]
    lines += [
        "",
        "## 9. Proposal skeleton",
    ]
    lines += [f"- {row.section_order}. {row.section_title}: {row.purpose or 'Purpose pending'} Inputs: {row.required_inputs or 'Pending Input'}" for row in skeleton] or ["- Pending Input: proposal skeleton not created."]
    lines += [
        "",
        "## 10. Methodology outline",
    ]
    lines += [
        f"- {row.module_order}. {row.module_title}: {row.module_description or 'Description pending'}\n  - Analysis approach: {row.analysis_approach or 'Pending Input'}\n  - Data collection conditions: {row.data_collection_conditions or 'Pending Input'}\n  - Evidence needed: {row.evidence_needed or 'Pending Input'}"
        for row in methodology
    ] or ["- Pending Input: methodology outline not created."]
    lines += [
        "",
        "## 11. Risk register",
    ]
    lines += [f"- [{risk.risk_id}] {risk.severity}/{risk.likelihood}: {risk.risk_description} Mitigation: {risk.mitigation or 'Pending Input'}" for risk in risks] or ["- No risks recorded."]
    lines += [
        "",
        "## 12. Clarification questions",
    ]
    lines += [f"- {q.question} Why it matters: {q.why_it_matters or 'Pending Input'}" for q in questions] or ["- No clarification questions recorded."]
    lines += [
        "",
        "## 13. Recommended next actions",
        "- Resolve all Gap and Pending Input rows before client-ready delivery.",
        "- Confirm evidence for mandatory requirements.",
        "- Review CV match language against provided evidence only.",
        "- Run verification before export.",
        "",
        "## 14. Upsell path",
        "- Full technical proposal sprint",
        "- CV rewrite support",
        "- Strategic bid partner package",
        "- Monthly bid desk retainer",
        "",
        "## 15. Bangladesh e-GP readiness",
    ]
    if egp_checks:
        lines += [f"- [{item.status}] {item.item_title}: {item.evidence_notes or 'No evidence note recorded.'}" for item in egp_checks]
    else:
        lines.append("- Not activated for this case. For Bangladesh e-GP procurements, review the live tender and current BPPA instructions before submission.")
    lines += [
        "",
        "## 16. Evidence reviews and approval record",
    ]
    reviews = session.query(EvidenceReview).filter_by(case_id=case.case_id).order_by(EvidenceReview.id).all()
    lines += [f"- [{row.evidence_id}] Requirement: {row.requirement_id}; compliance: {row.compliance_id}; state: {row.review_status}; entity: {row.legal_entity}; reviewer: {row.reviewed_by}; expiry: {row.valid_until or row.validity_reason}; source: {row.source_id or 'archived attachment'}; evidence SHA-256: {row.file_sha256 or row.evidence_text_sha256}." for row in reviews] or ["- Pending Input: no attributable evidence reviews recorded."]
    lines += [
        "",
        "Prepared for internal/client review.",
    ]
    report = "\n".join(lines)
    return _remove_prohibited_phrases(report)


def _record_output(session: Session, case_id: str, output_type: str, path: Path, context) -> OutputPackage:
    from .verification_service import _scan_outputs_for_prohibited_phrases

    try:
        validate_publication(session, case_id, context)
        issues = _scan_outputs_for_prohibited_phrases([OutputPackage(output_type=output_type, file_path=str(path))])
        if issues:
            raise ValueError("Report content could not pass safety inspection: " + "; ".join(issues))
    except Exception:
        quarantine_artifact(path)
        raise
    published = path.with_name(path.name.replace(".pending.", "."))
    path.rename(published)
    path = published
    output = OutputPackage(
            output_id=f"OUT-{uuid4().hex[:10].upper()}",
            case_id=case_id,
            output_type=output_type,
            file_path=str(path),
            readiness_label=context["readiness_label"],
            verification_status="approved" if context["approved"] else "draft",
            snapshot_hash=context["snapshot_hash"],
            approval_id=context["approval"].approval_id if context["approved"] else None,
            verification_run_id=context["run"].run_id if context["run"] else None,
            content_sha256=file_hash(path),
        )
    session.add(output)
    case = session.query(Case).filter_by(case_id=case_id).one()
    case.readiness_label = context["readiness_label"]
    case.status = "packaged" if context["readiness_label"] == "EXPORT_READY" else "client_ready" if context["readiness_label"] == "CLIENT_READY" else "blocked" if context["blockers"] else "preliminary"
    session.flush()
    return output


def _remove_prohibited_phrases(text: str) -> str:
    import re
    cleaned = text
    for phrase in PROHIBITED_REPORT_PHRASES:
        cleaned = re.sub(re.escape(phrase), "[removed prohibited claim]", cleaned, flags=re.IGNORECASE)
    return cleaned
