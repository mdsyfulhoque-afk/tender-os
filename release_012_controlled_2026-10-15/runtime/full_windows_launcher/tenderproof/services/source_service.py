from pathlib import Path
from uuid import uuid4
import json

from docx import Document
from pypdf import PdfReader
from sqlalchemy.orm import Session

from ..config import EXTRACTED_TEXT_DIR, UPLOADS_DIR, ensure_directories
from ..models import Source
from .audit_service import log_event
from .tor_normalizer import create_gap_for_failed_extraction, normalize_text
from .decision_service import file_hash, invalidate_case


def _version_source(session, source):
    previous = session.query(Source).filter_by(case_id=source.case_id, filename=source.filename).filter(Source.source_id != source.source_id).order_by(Source.id.desc()).first()
    source.version_number = ((previous.version_number or 1) + 1) if previous else 1
    source.supersedes_source_id = previous.source_id if previous else None
    source.content_sha256 = file_hash(source.file_path) if source.file_path else file_hash(source.text_extract_path)
    source.text_sha256 = file_hash(source.text_extract_path)
    invalidate_case(session, source.case_id, f"Registered source version {source.version_number}: {source.source_id}")


def register_manual_text(session: Session, case_id: str, text: str, filename: str = "manual_tor.txt") -> Source:
    ensure_directories()
    source_id = f"SRC-{uuid4().hex[:10].upper()}"
    extract_path = EXTRACTED_TEXT_DIR / f"{source_id}.txt"
    raw_path = UPLOADS_DIR / f"{source_id}_raw.txt"
    raw_path.write_bytes(text.encode("utf-8"))
    extract_path.write_text(normalize_text(text), encoding="utf-8")
    source = Source(
        source_id=source_id,
        case_id=case_id,
        filename=filename,
        file_type="txt",
        source_category="Manual ToR Text",
        file_path=str(raw_path),
        text_extract_path=str(extract_path),
        extraction_status="extracted" if text.strip() else "empty",
        notes="Manual ToR text pasted into app.",
        extraction_coverage_json=json.dumps({"method": "manual", "complete": bool(text.strip()), "items": 1}),
    )
    session.add(source)
    session.flush()
    _version_source(session, source)
    log_event(session, case_id, "source.manual_text", f"Registered manual ToR text as {source_id}")
    if not text.strip():
        create_gap_for_failed_extraction(session, case_id, source_id, "Manual text is empty.")
    return source


def save_uploaded_file(session: Session, case_id: str, uploaded_file, source_category: str = "ToR", progress_callback=None) -> Source:
    ensure_directories()
    source_id = f"SRC-{uuid4().hex[:10].upper()}"
    suffix = Path(uploaded_file.name).suffix.lower()
    file_path = UPLOADS_DIR / f"{source_id}{suffix}"
    file_path.write_bytes(uploaded_file.getbuffer())
    coverage = {}
    text, status, notes = extract_text(file_path, progress_callback=progress_callback, coverage=coverage)
    extract_path = EXTRACTED_TEXT_DIR / f"{source_id}.txt"
    extract_path.write_text(text, encoding="utf-8")
    source = Source(
        source_id=source_id,
        case_id=case_id,
        filename=uploaded_file.name,
        file_type=suffix.lstrip(".") or "unknown",
        source_category=source_category,
        file_path=str(file_path),
        text_extract_path=str(extract_path),
        extraction_status=status,
        notes=notes,
        extraction_coverage_json=json.dumps(coverage, sort_keys=True),
    )
    session.add(source)
    session.flush()
    _version_source(session, source)
    if status != "extracted":
        create_gap_for_failed_extraction(session, case_id, source_id, notes)
    log_event(session, case_id, "source.uploaded", f"Uploaded source {uploaded_file.name}", {"source_id": source_id})
    return source


def extract_text(file_path: Path, progress_callback=None, coverage=None) -> tuple[str, str, str]:
    coverage = coverage if coverage is not None else {}
    suffix = file_path.suffix.lower()
    try:
        if suffix == ".txt":
            text = normalize_text(file_path.read_text(encoding="utf-8-sig"))
            coverage.update(method="text", complete=bool(text))
            return text, "extracted" if text else "empty", "TXT extracted." if text else "TXT is empty."
        if suffix == ".docx":
            doc = Document(str(file_path))
            text = "\n".join([p.text for p in doc.paragraphs] + [" | ".join(cell.text for cell in row.cells) for table in doc.tables for row in table.rows])
            coverage.update(method="docx", complete=bool(text.strip()), tables=len(doc.tables), paragraphs=len(doc.paragraphs))
            return normalize_text(text), "extracted" if text.strip() else "empty", "DOCX extracted."
        if suffix == ".pdf":
            reader = PdfReader(str(file_path))
            pages = [page.extract_text() or "" for page in reader.pages]
            text = "\n".join(pages)
            coverage.update(method="pdf", total_pages=len(pages), text_layer_pages=[i + 1 for i, page in enumerate(pages) if len(page.strip()) >= 20])
            if pages and all(len(page.strip()) >= 20 for page in pages):
                coverage.update(complete=True, processed_pages=len(pages), unreadable_pages=[])
                return normalize_text(text), "extracted", "PDF extracted with pypdf."
            ocr_text, ocr_note = _ocr_pdf(file_path, progress_callback=progress_callback, coverage=coverage, text_pages=pages)
            if ocr_text.strip():
                status = "extracted" if len(ocr_text.strip()) >= 40 and coverage.get("complete") else "partial"
                return normalize_text(ocr_text), status, ocr_note
            return normalize_text(text), "empty", ocr_note
        return "", "unsupported", f"Unsupported file type: {suffix}"
    except Exception as exc:
        coverage.update(complete=False, error=type(exc).__name__)
        return "", "failed", str(exc)


def _ocr_pdf(file_path: Path, progress_callback=None, max_pages: int = 200, coverage=None, text_pages=None) -> tuple[str, str]:
    """OCR image-only PDFs when the optional local OCR dependencies are installed."""
    coverage = coverage if coverage is not None else {}
    coverage.update(method="pdf_mixed_ocr", complete=False)
    try:
        import numpy as np
        import pypdfium2 as pdfium
        from rapidocr import RapidOCR
    except ImportError:
        return "", "No selectable PDF text found. OCR unavailable; install project OCR requirements to process scanned PDFs."

    try:
        engine = RapidOCR()
        document = pdfium.PdfDocument(str(file_path))
        page_limit = min(len(document), max_pages)
        output = []
        unreadable = []
        for index in range(page_limit):
            if text_pages and len(text_pages[index].strip()) >= 20:
                page_text = text_pages[index]
            else:
                image = document[index].render(scale=2).to_pil()
                result = engine(np.asarray(image.convert("RGB")))
                page_text = "\n".join(text for text in (result.txts or []) if str(text).strip())
            if page_text:
                output.append(f"[OCR page {index + 1}]\n{page_text}")
            else:
                unreadable.append(index + 1)
            if progress_callback:
                progress_callback(index + 1, page_limit)
        coverage.update(total_pages=len(document), processed_pages=page_limit, unreadable_pages=unreadable, complete=page_limit == len(document) and not unreadable and bool(output))
        total_pages = len(document)
        close = getattr(document, "close", None)
        if close:
            close()
        if not output:
            return "", "OCR ran but found no readable text in the scanned PDF."
        note = "Scanned PDF processed with local RapidOCR. Review extracted text against the source."
        if total_pages > page_limit:
            note += f" OCR limited to first {page_limit} of {total_pages} pages."
        if unreadable:
            note += f" OCR unreadable pages: {unreadable}."
        return "\n\n".join(output), note
    except Exception as exc:
        return "", f"OCR failed: {exc}"


def get_source_text(source: Source) -> str:
    if source.text_extract_path and Path(source.text_extract_path).exists():
        return Path(source.text_extract_path).read_text(encoding="utf-8", errors="ignore")
    return ""


def list_sources(session: Session, case_id: str) -> list[Source]:
    return session.query(Source).filter(Source.case_id == case_id).order_by(Source.uploaded_at.desc()).all()

