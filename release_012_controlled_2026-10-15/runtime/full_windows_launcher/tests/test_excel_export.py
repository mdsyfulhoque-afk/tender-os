from openpyxl import load_workbook

from tenderproof.services.export_service import REQUIRED_SHEETS, export_case_to_excel
from tenderproof.services.sample_data_service import create_demo_case
from tenderproof.services.egp_service import ensure_egp_checklist


def test_workbook_is_created(session):
    case = create_demo_case(session)
    path = export_case_to_excel(session, case.case_id)
    assert path.exists()


def test_all_required_sheets_exist(session):
    case = create_demo_case(session)
    ensure_egp_checklist(session, case.case_id)
    path = export_case_to_excel(session, case.case_id)
    workbook = load_workbook(path)
    assert set(REQUIRED_SHEETS).issubset(set(workbook.sheetnames))

