import json
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from datetime import datetime
from pathlib import Path
from app.config import DATA_DIR

def generate_proposal(tender_id: str):
    tender_dir = DATA_DIR / tender_id
    
    with open(tender_dir / "client_submission.json", "r", encoding="utf-8") as f:
        submission = json.load(f)
    with open(tender_dir / "doc_checklist.json", "r", encoding="utf-8") as f:
        checklist = json.load(f)
    with open(tender_dir / "proposal_keywords.json", "r", encoding="utf-8") as f:
        keywords = json.load(f)
        
    doc = Document()
    title = doc.add_heading('TENDER SUBMISSION LETTER', 0)
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    
    doc.add_paragraph(f"Date: {datetime.now().strftime('%B %d, %Y')}")
    doc.add_paragraph("To,\nThe Procuring Entity,\n[Procuring Agency Name],\n[Address]")
    doc.add_paragraph("Subject: Submission of Tender for [Tender Name/Reference]").add_run().bold = True
    doc.add_paragraph("Dear Sir/Madam,")
    
    kw_str = ", ".join([k["keyword"] for k in keywords[:5]])
    doc.add_paragraph(
        f"With reference to the above subject, we are pleased to submit our technical and financial proposal. "
        f"Our firm is fully committed to delivering the project with a focus on {kw_str}, ensuring compliance with all technical specifications and scope of work outlined in the tender document."
    )
    
    doc.add_paragraph("We confirm that our company meets all the eligibility and financial criteria required for this tender:")
    for key, value in submission.items():
        p = doc.add_paragraph(style='List Bullet')
        p.add_run(f"{key.replace('_', ' ').title()}: ").bold = True
        p.add_run(str(value))
        
    doc.add_paragraph("Furthermore, we have enclosed all the mandatory documents as per the tender checklist:")
    for item in checklist:
        if item.get("mandatory"):
            p = doc.add_paragraph(style='List Bullet')
            p.add_run(f"{item['name']} (Ref: {item['clause_reference']})")
            
    doc.add_paragraph("We assure you of our best services at all times and look forward to a positive response.")
    doc.add_paragraph("Yours faithfully,\n[Authorized Signatory Name]\n[Designation]\n[Company Name]")
    
    doc.save(str(tender_dir / "generated_proposal.docx"))