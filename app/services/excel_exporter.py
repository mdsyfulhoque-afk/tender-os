import json
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from pathlib import Path
from app.config import DATA_DIR

def generate_excel(tender_id: str):
    tender_dir = DATA_DIR / tender_id
    
    with open(tender_dir / "doc_checklist.json", "r", encoding="utf-8") as f:
        checklist = json.load(f)
    with open(tender_dir / "proposal_keywords.json", "r", encoding="utf-8") as f:
        keywords = json.load(f)
    with open(tender_dir / "client_form_schema.json", "r", encoding="utf-8") as f:
        form_schema = json.load(f)
        
    wb = openpyxl.Workbook()
    header_fill = PatternFill(start_color="1F4E79", end_color="1F4E79", fill_type="solid")
    header_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
    border = Border(left=Side(style='thin'), right=Side(style='thin'), top=Side(style='thin'), bottom=Side(style='thin'))
    
    def style_header(ws):
        for row in ws[1]:
            row.fill = header_fill
            row.font = header_font
            row.alignment = Alignment(horizontal='center', vertical='center')
            row.border = border
            
    def auto_width(ws):
        for col in ws.columns:
            max_length = 0
            column = col[0].column_letter
            for cell in col:
                try:
                    if len(str(cell.value)) > max_length:
                        max_length = len(str(cell.value))
                except:
                    pass
            ws.column_dimensions[column].width = (max_length + 2) * 1.2

    ws_check = wb.active
    ws_check.title = "Document Checklist"
    ws_check.append(["Document Name", "Clause Reference", "Mandatory", "Description"])
    style_header(ws_check)
    for doc in checklist:
        ws_check.append([doc["name"], doc["clause_reference"], "YES" if doc["mandatory"] else "NO", doc["description"]])
    auto_width(ws_check)
    
    ws_key = wb.create_sheet("Proposal Keywords")
    ws_key.append(["Keyword / Term", "Recommended Section", "Context Excerpt"])
    style_header(ws_key)
    for k in keywords:
        ws_key.append([k["keyword"], k["section"], k["context"]])
    auto_width(ws_key)
    
    ws_form = wb.create_sheet("Client Form Requirements")
    ws_form.append(["Field Label", "Input Type", "Required", "Description"])
    style_header(ws_form)
    for f in form_schema:
        ws_form.append([f["label"], f["input_type"], "YES" if f["required"] else "NO", f["description"]])
    auto_width(ws_form)
    
    wb.save(str(tender_dir / "tender_requirements.xlsx"))