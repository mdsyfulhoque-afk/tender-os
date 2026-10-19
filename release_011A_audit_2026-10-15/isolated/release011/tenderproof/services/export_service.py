from datetime import UTC, datetime
from pathlib import Path
from uuid import uuid4

import pandas as pd
from openpyxl.styles import Font, Alignment
from openpyxl.utils import get_column_letter
from sqlalchemy.orm import Session

from ..config import EXPORTS_DIR, ensure_directories
from ..models import (
    AuditEvent,
    BidRecommendation,
    BidScorecard,
    Case,
    ClarificationQuestion,
    ComplianceLedger,
    CVMatch,
    EGPChecklistItem,
    GapItem,
    IPCapture,
    MethodologyOutline,
    NormalizedRequirement,
    OutputPackage,
    ProposalSkeleton,
    RiskRegister,
    Source,
    VerificationResult,
)
from .audit_service import log_event


SHEET_MODELS = [
    ("Case Summary", Case),
    ("Source Files", Source),
    ("ToR Extraction", NormalizedRequirement),
    ("Compliance Matrix", ComplianceLedger),
    ("CV Matcher", CVMatch),
    ("Bid No-Bid Scorecard", BidScorecard),
    ("Bid Recommendation", BidRecommendation),
    ("Risk Register", RiskRegister),
    ("Gap Items", GapItem),
    ("Clarification Questions", ClarificationQuestion),
    ("Proposal Skeleton", ProposalSkeleton),
    ("Methodology Outline", MethodologyOutline),
    ("IP Capture", IPCapture),
    ("Verification Results", VerificationResult),
    ("Audit Log", AuditEvent),
    ("e-GP Readiness", EGPChecklistItem),
]


REQUIRED_SHEETS = [name for name, _ in SHEET_MODELS]


def export_case_to_excel(session: Session, case_id: str) -> Path:
    ensure_directories()
    path = EXPORTS_DIR / f"{case_id}_TenderProof_Diagnostic_Workbook.xlsx"
    with pd.ExcelWriter(path, engine="openpyxl") as writer:
        for sheet_name, model in SHEET_MODELS:
            query = session.query(model)
            if hasattr(model, "case_id"):
                query = query.filter(model.case_id == case_id)
            rows = [_model_to_dict(row) for row in query.all()]
            if sheet_name == "Case Summary":
                rows = rows or [{"case_id": case_id}]
                rows[0]["exported_at_utc"] = datetime.now(UTC).isoformat(timespec="seconds")
            pd.DataFrame(rows).to_excel(writer, sheet_name=sheet_name[:31], index=False)
        workbook = writer.book
        for sheet in workbook.worksheets:
            sheet.freeze_panes = "A2"
            for cell in sheet[1]:
                cell.font = Font(bold=True)
                cell.alignment = Alignment(wrap_text=True, vertical="top")
            for row in sheet.iter_rows(min_row=2):
                for cell in row:
                    cell.alignment = Alignment(wrap_text=True, vertical="top")
            for idx, column_cells in enumerate(sheet.columns, start=1):
                values = [str(cell.value or "") for cell in column_cells]
                width = min(max(max((len(v) for v in values), default=8) + 2, 12), 55)
                sheet.column_dimensions[get_column_letter(idx)].width = width
    _record_export(session, case_id, path)
    log_event(session, case_id, "export.xlsx", f"Generated Excel workbook {path.name}")
    return path


def _model_to_dict(row) -> dict:
    data = {}
    for column in row.__table__.columns:
        value = getattr(row, column.name)
        data[column.name] = value.isoformat(timespec="seconds") if hasattr(value, "isoformat") else value
    return data


def _record_export(session: Session, case_id: str, path: Path) -> OutputPackage:
    existing = session.query(OutputPackage).filter(OutputPackage.case_id == case_id, OutputPackage.output_type == "excel_export").first()
    if existing:
        existing.file_path = str(path)
        existing.verification_status = "not_run"
        output = existing
    else:
        output = OutputPackage(
            output_id=f"OUT-{uuid4().hex[:10].upper()}",
            case_id=case_id,
            output_type="excel_export",
            file_path=str(path),
            readiness_label="DRAFT_ONLY",
            verification_status="not_run",
        )
        session.add(output)
    session.flush()
    return output
