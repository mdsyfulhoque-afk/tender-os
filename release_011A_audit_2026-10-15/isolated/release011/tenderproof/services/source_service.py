from pathlib import Path
from uuid import uuid4

from docx import Document
from pypdf import PdfReader
from sqlalchemy.orm import Session

from ..config import EXTRACTED_TEXT_DIR, UPLOADS_DIR, ensure_directories
from ..models import Source, Case, OutputPackage, VerificationResult
from .audit_service import log_event
from .tor_normalizer import create_gap_for_failed_extraction, normalize_text



def _invalidate_readiness(session: Session, case_id: str) -> None:
    case = session.query(Case).filter_by(case_id=case_id).first()
    if case is None:
        raise ValueError(f"Unknown case: {case_id}")

    case.readiness_label = "DRAFT_ONLY"
    case.status = "draft"

    for output in session.query(OutputPackage).filter_by(
        case_id=case_id
    ).all():
        output.readiness_label = "DRAFT_ONLY"
        output.verification_status = "stale"

    session.query(VerificationResult).filter_by(
        case_id=case_id
    ).delete(synchronize_session=False)

    log_event(
        session, case_id, "verification.invalidated",
        "New source registered; prior readiness requires re-verification."
    )

def register_manual_text(session: Session, case_id: str, text: str, filename: str = "manual_tor.txt") -> Source:
    ensure_directories()
    source_id = f"SRC-{uuid4().hex[:10].upper()}"
    extract_path = EXTRACTED_TEXT_DIR / f"{source_id}.txt"
    extract_path.write_text(normalize_text(text), encoding="utf-8")
    source = Source(
        source_id=source_id,
        case_id=case_id,
        filename=filename,
        file_type="txt",
        source_category="Manual ToR Text",
        file_path="",
        text_extract_path=str(extract_path),
        extraction_status="extracted" if text.strip() else "empty",
        notes="Manual ToR text pasted into app.",
    )
    session.add(source)
    session.flush()
    _invalidate_readiness(session, case_id)
    log_event(session, case_id, "source.manual_text", f"Registered manual ToR text as {source_id}")
    return source


def save_uploaded_file(session: Session, case_id: str, uploaded_file, source_category: str = "ToR", progress_callback=None) -> Source:
    ensure_directories()
    source_id = f"SRC-{uuid4().hex[:10].upper()}"
    suffix = Path(uploaded_file.name).suffix.lower()
    file_path = UPLOADS_DIR / f"{source_id}{suffix}"
    file_path.write_bytes(uploaded_file.getbuffer())
    text, status, notes = extract_text(file_path, progress_callback=progress_callback)
    extract_path = EXTRACTED_TEXT_DIR / f"{source_id}.txt"
    extract_path.write_text(text, encoding="utf-8")
    source = Source(
        source_id=source_id,
        case_id=case_id,
        filename=uploaded_file.name,
        file_type=suffix.lstrip(".") or "unknown",
        source_category=source_category,
        file_path=str(file_path),
        text_extract_path=str(extract_path),
        extraction_status=status,
        notes=notes,
    )
    session.add(source)
    session.flush()
    _invalidate_readiness(session, case_id)
    if status != "extracted":
        create_gap_for_failed_extraction(session, case_id, source_id, notes)
    log_event(session, case_id, "source.uploaded", f"Uploaded source {uploaded_file.name}", {"source_id": source_id})
    return source


def extract_text(file_path: Path, progress_callback=None) -> tuple[str, str, str]:
    suffix = file_path.suffix.lower()
    try:
        if suffix == ".txt":
            return normalize_text(file_path.read_text(encoding="utf-8", errors="ignore")), "extracted", "TXT extracted."
        if suffix == ".docx":
            doc = Document(str(file_path))
            text = "\n".join(p.text for p in doc.paragraphs)
            return normalize_text(text), "extracted" if text.strip() else "empty", "DOCX extracted."
        if suffix == ".pdf":
            reader = PdfReader(str(file_path))
            pages = [page.extract_text() or "" for page in reader.pages]
            text = "\n".join(pages)
            if sum(len(page.strip()) for page in pages) >= max(40, len(pages) * 20):
                return normalize_text(text), "extracted", "PDF extracted with pypdf."
            ocr_text, ocr_note = _ocr_pdf(file_path, progress_callback=progress_callback)
            if ocr_text.strip():
                status = "extracted" if len(ocr_text.strip()) >= 40 else "partial"
                return normalize_text(ocr_text), status, ocr_note
            return normalize_text(text), "empty", ocr_note
        return "", "unsupported", f"Unsupported file type: {suffix}"
    except Exception as exc:
        return "", "failed", str(exc)


def _ocr_pdf(file_path: Path, progress_callback=None, max_pages: int = 200) -> tuple[str, str]:
    """OCR image-only PDFs when the optional local OCR dependencies are installed."""
    try:
        import numpy as np
        import pypdfium2 as pdfium
        from rapidocr import RapidOCR
    except ImportError:
        return "", "No selectable PDF text found. OCR unavailable; install project OCR requirements to process scanned PDFs."

    try:
        engine = RapidOCR()
        document = pdfium.PdfDocument(str(file_path))
        page_limit = min(len(document), max_pages)
        output = []
        for index in range(page_limit):
            image = document[index].render(scale=2).to_pil()
            result = engine(np.asarray(image.convert("RGB")))
            page_text = "\n".join(text for text in (result.txts or []) if str(text).strip())
            if page_text:
                output.append(f"[OCR page {index + 1}]\n{page_text}")
            if progress_callback:
                progress_callback(index + 1, page_limit)
        if not output:
            return "", "OCR ran but found no readable text in the scanned PDF."
        note = "Scanned PDF processed with local RapidOCR. Review extracted text against the source."
        if len(document) > page_limit:
            note += f" OCR limited to first {page_limit} of {len(document)} pages."
        return "\n\n".join(output), note
    except Exception as exc:
        return "", f"OCR failed: {exc}"


def get_source_text(source: Source) -> str:
    if source.text_extract_path and Path(source.text_extract_path).exists():
        return Path(source.text_extract_path).read_text(encoding="utf-8", errors="ignore")
    return ""


def list_sources(session: Session, case_id: str) -> list[Source]:
    return session.query(Source).filter(Source.case_id == case_id).order_by(Source.uploaded_at.desc()).all()

