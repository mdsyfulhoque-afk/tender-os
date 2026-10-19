from datetime import UTC, datetime
from pathlib import Path
from uuid import uuid4

import pandas as pd
from openpyxl.styles import Font, Alignment
from openpyxl.utils import get_column_letter
from sqlalchemy.orm import Session

from ..config import EXPORTS_DIR, ensure_directories
from ..models import (
    AuditEvent,
    BidRecommendation,
    BidScorecard,
    Case,
    ClarificationQuestion,
    ComplianceLedger,
    CVMatch,
    EGPChecklistItem,
    GapItem,
    IPCapture,
    MethodologyOutline,
    NormalizedRequirement,
    OutputPackage,
    ProposalSkeleton,
    RiskRegister,
    Source,
    VerificationResult,
)
from .audit_service import log_event
from .decision_service import file_hash, release_context, snapshot_hash


SHEET_MODELS = [
    ("Case Summary", Case),
    ("Source Files", Source),
    ("ToR Extraction", NormalizedRequirement),
    ("Compliance Matrix", ComplianceLedger),
    ("CV Matcher", CVMatch),
    ("Bid No-Bid Scorecard", BidScorecard),
    ("Bid Recommendation", BidRecommendation),
    ("Risk Register", RiskRegister),
    ("Gap Items", GapItem),
    ("Clarification Questions", ClarificationQuestion),
    ("Proposal Skeleton", ProposalSkeleton),
    ("Methodology Outline", MethodologyOutline),
    ("IP Capture", IPCapture),
    ("Verification Results", VerificationResult),
    ("Audit Log", AuditEvent),
    ("e-GP Readiness", EGPChecklistItem),
]


REQUIRED_SHEETS = [name for name, _ in SHEET_MODELS]


def export_case_to_excel(session: Session, case_id: str) -> Path:
    ensure_directories()
    if session.query(Case).filter_by(case_id=case_id).first() is None:
        raise ValueError("Case not found.")
    context = release_context(session, case_id, generating="excel")
    path = EXPORTS_DIR / f"{case_id}_TenderProof_Diagnostic_Workbook_{uuid4().hex[:10]}.xlsx"
    with pd.ExcelWriter(path, engine="openpyxl") as writer:
        for sheet_name, model in SHEET_MODELS:
            query = session.query(model)
            if hasattr(model, "case_id"):
                query = query.filter(model.case_id == case_id)
            rows = [_model_to_dict(row) for row in query.all()]
            if sheet_name == "Case Summary":
                rows = rows or [{"case_id": case_id}]
                rows[0]["exported_at_utc"] = datetime.now(UTC).isoformat(timespec="seconds")
                rows[0]["readiness_label"] = context["readiness_label"]
                rows[0]["status"] = {"BLOCKED": "blocked", "CLIENT_READY": "client_ready", "EXPORT_READY": "packaged"}.get(context["readiness_label"], "preliminary")
                rows[0]["snapshot_hash"] = context["snapshot_hash"]
                rows[0]["eligibility"] = context["eligibility"]
                rows[0]["approval_id"] = context["approval"].approval_id if context["approved"] else "Not approved"
                rows[0]["approved_by"] = context["approval"].approved_by if context["approved"] else ""
                rows[0]["release_status"] = "APPROVED" if context["approved"] and context["readiness_label"] == "EXPORT_READY" else "DRAFT"
            if sheet_name == "Bid Recommendation":
                for row in rows:
                    row["provisional_opportunity_score_only"] = True
                    row["current_eligibility"] = context["eligibility"]
                    row["current_release_clearance"] = "Approved" if context["approved"] else "No clearance"
            if sheet_name == "Verification Results":
                for row in rows:
                    row["current_snapshot_run"] = bool(context["run"] and row.get("run_id") == context["run"].run_id)
            pd.DataFrame(rows).to_excel(writer, sheet_name=sheet_name[:31], index=False)
        workbook = writer.book
        for sheet in workbook.worksheets:
            sheet.freeze_panes = "A2"
            for cell in sheet[1]:
                cell.font = Font(bold=True)
                cell.alignment = Alignment(wrap_text=True, vertical="top")
            for row in sheet.iter_rows(min_row=2):
                for cell in row:
                    if cell.data_type == "f":
                        cell.data_type = "s"
                    cell.alignment = Alignment(wrap_text=True, vertical="top")
            for idx, column_cells in enumerate(sheet.columns, start=1):
                values = [str(cell.value or "") for cell in column_cells]
                width = min(max(max((len(v) for v in values), default=8) + 2, 12), 55)
                sheet.column_dimensions[get_column_letter(idx)].width = width
    if snapshot_hash(session, case_id) != context["snapshot_hash"]:
        raise ValueError("Assessment changed during generation; regenerate the workbook.")
    _record_export(session, case_id, path, context)
    log_event(session, case_id, "export.xlsx", f"Generated Excel workbook {path.name}")
    return path


def _model_to_dict(row) -> dict:
    data = {}
    for column in row.__table__.columns:
        value = getattr(row, column.name)
        data[column.name] = value.isoformat(timespec="seconds") if hasattr(value, "isoformat") else value
    return data


def _record_export(session: Session, case_id: str, path: Path, context) -> OutputPackage:
    output = OutputPackage(
            output_id=f"OUT-{uuid4().hex[:10].upper()}",
            case_id=case_id,
            output_type="excel_export",
            file_path=str(path),
            readiness_label=context["readiness_label"],
            verification_status="approved" if context["approved"] and context["readiness_label"] == "EXPORT_READY" else "draft",
            snapshot_hash=context["snapshot_hash"],
            approval_id=context["approval"].approval_id if context["approved"] and context["readiness_label"] == "EXPORT_READY" else None,
            verification_run_id=context["run"].run_id if context["run"] else None,
            content_sha256=file_hash(path),
        )
    session.add(output)
    case = session.query(Case).filter_by(case_id=case_id).one()
    case.readiness_label = context["readiness_label"]
    case.status = "packaged" if context["readiness_label"] == "EXPORT_READY" else "blocked" if context["blockers"] else "preliminary"
    session.flush()
    return output
