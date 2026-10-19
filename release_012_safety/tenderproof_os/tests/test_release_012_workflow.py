from datetime import UTC, datetime, timedelta
from pathlib import Path
import json

import pytest
from openpyxl import load_workbook

from tenderproof.constants import REQUIRED_BID_CATEGORIES
from tenderproof.models import DecisionApproval, EvidenceReview, OutputPackage, Source, VerificationResult, VerificationRun
from tenderproof.services.approval_service import approve_case, record_evidence_review, review_source, revoke_approval
from tenderproof.services.bid_score_service import save_scores_and_recommendation
from tenderproof.services.case_service import create_case
from tenderproof.services.compliance_service import auto_create_compliance_rows, update_compliance_row
from tenderproof.services.decision_service import release_context
from tenderproof.services.export_service import export_case_to_excel
from tenderproof.services.extraction_service import add_manual_requirement, update_requirement
from tenderproof.services.report_service import generate_docx_report, generate_markdown_report
from tenderproof.services.source_service import register_manual_text
from tenderproof.services.verification_service import run_verification


def reviewed_assessment(session, tmp_path):
    case = create_case(session, client_name="Synthetic Verified Company", organization="Synthetic Verified Company")
    source = register_manual_text(session, case.case_id, "Bidder must hold a current license.", filename="tender.txt")
    requirement = add_manual_requirement(session, case.case_id, source_id=source.source_id, requirement_text="Bidder must hold a current license.", mandatory=True, needs_review=False, tor_section_page="Clause 1")
    row = auto_create_compliance_rows(session, case.case_id)[0]
    update_compliance_row(session, row.compliance_id, status="Met", available_evidence="License ABC issued to Synthetic Verified Company.")
    review_source(session, case.case_id, source.source_id, reviewer="Synthetic QA reviewer", notes="Checked source and all applicable clause coverage.", requirements_complete=True)
    proof = tmp_path / "license.txt"
    proof.write_text("Synthetic license ABC for Synthetic Verified Company.", encoding="utf-8")
    evidence = record_evidence_review(session, case.case_id, row.compliance_id, reviewer="Synthetic QA reviewer", notes="Checked original license, bidder identity and validity.", legal_entity=case.organization, file_path=proof, valid_until=(datetime.now(UTC).date() + timedelta(days=30)).isoformat())
    scores = {c: {"score": 5, "rationale": "Synthetic fit", "evidence_or_gap_reference": evidence.evidence_id} for c in REQUIRED_BID_CATEGORIES}
    save_scores_and_recommendation(session, case.case_id, scores)
    return case, source, requirement, row, evidence, proof


def approved_assessment(session, tmp_path):
    values = reviewed_assessment(session, tmp_path)
    case = values[0]
    result = run_verification(session, case.case_id)
    assert result.overall_status == "warning"
    assert all(r.rule_code == "APP-001" for r in result.warnings)
    approval = approve_case(session, case.case_id, reviewer="Synthetic founder", rationale="Reviewed this exact synthetic assessment for release.")
    return (*values, approval)


def test_positive_review_approve_report_and_workbook_workflow(session, tmp_path):
    case, *_, approval = approved_assessment(session, tmp_path)
    report = generate_markdown_report(session, case.case_id)
    assert "Readiness label: **CLIENT_READY**" in report.read_text(encoding="utf-8")
    assert approval.approval_id in report.read_text(encoding="utf-8")
    assert generate_docx_report(session, case.case_id).is_file()
    workbook = load_workbook(export_case_to_excel(session, case.case_id))
    summary = workbook["Case Summary"]
    values = dict(zip([c.value for c in summary[1]], [c.value for c in summary[2]]))
    assert values["readiness_label"] == "EXPORT_READY"
    assert values["approval_id"] == approval.approval_id
    assert len(workbook.sheetnames) == 16
    verified = run_verification(session, case.case_id)
    assert verified.overall_status == "passed"
    assert verified.readiness_label == "EXPORT_READY"


def test_approval_requires_current_verification_and_named_human(session, tmp_path):
    case, *_ = reviewed_assessment(session, tmp_path)
    with pytest.raises(ValueError):
        approve_case(session, case.case_id, reviewer="Reviewer", rationale="No verification yet.")
    run_verification(session, case.case_id)
    with pytest.raises(ValueError):
        approve_case(session, case.case_id, reviewer="", rationale="No attribution.")


def test_approval_and_verification_history_survive_amendment(session, tmp_path):
    case, source, *_, approval = approved_assessment(session, tmp_path)
    generated = generate_markdown_report(session, case.case_id)
    previous_bytes = generated.read_bytes()
    run_count = session.query(VerificationRun).count()
    rule_count = session.query(VerificationResult).count()
    amended = register_manual_text(session, case.case_id, "Bidder must also hold a new certificate.", filename="tender.txt")
    assert amended.version_number == 2
    assert amended.supersedes_source_id == source.source_id
    assert session.query(DecisionApproval).count() == 1
    assert approval.approved_by == "Synthetic founder"
    assert run_verification(session, case.case_id).readiness_label == "BLOCKED"
    assert session.query(VerificationRun).count() == run_count + 1
    assert session.query(VerificationResult).count() > rule_count
    assert generated.read_bytes() == previous_bytes
    assert not release_context(session, case.case_id)["approved"]


@pytest.mark.parametrize("mutation", ["file_removed", "file_changed", "expired", "wrong_entity", "text_changed"])
def test_evidence_mutations_prevent_approved_exports(session, tmp_path, mutation):
    case, _, _, row, evidence, proof, _ = approved_assessment(session, tmp_path)
    generate_markdown_report(session, case.case_id)
    if mutation == "file_removed":
        proof.unlink()
    elif mutation == "file_changed":
        proof.write_text("Different evidence", encoding="utf-8")
    elif mutation == "expired":
        evidence.valid_until = (datetime.now(UTC).date() - timedelta(days=1)).isoformat()
    elif mutation == "wrong_entity":
        evidence.legal_entity = "Different Bidder"
    else:
        row.available_evidence = "A different claim."
    session.flush()
    context = release_context(session, case.case_id)
    assert not context["approved"]
    assert context["readiness_label"] not in {"CLIENT_READY", "EXPORT_READY"}
    export_case_to_excel(session, case.case_id)
    output = session.query(OutputPackage).filter_by(output_type="excel_export").order_by(OutputPackage.id.desc()).first()
    assert output.readiness_label not in {"CLIENT_READY", "EXPORT_READY"}
    assert output.approval_id is None
    with pytest.raises(ValueError):
        approve_case(session, case.case_id, reviewer="Reviewer", rationale="Attempt to reuse stale verification.")


def test_new_report_never_overwrites_historical_artifact(session, tmp_path):
    case, *_ = approved_assessment(session, tmp_path)
    first = generate_markdown_report(session, case.case_id)
    content = first.read_bytes()
    second = generate_markdown_report(session, case.case_id)
    assert first != second
    assert first.read_bytes() == content
    assert session.query(OutputPackage).filter_by(output_type="markdown_report").count() == 2


def test_requirement_original_quotation_is_preserved(session, tmp_path):
    case, _, requirement, *_ = approved_assessment(session, tmp_path)
    original = requirement.requirement_text
    update_requirement(session, requirement.requirement_id, requirement_text="Edited interpretation", atomic_requirement_text="New threshold", needs_review=True)
    assert requirement.requirement_text == original
    assert case.readiness_label == "DRAFT_ONLY"
    assert not release_context(session, case.case_id)["approved"]


def test_revoked_approval_cannot_release_new_artifact(session, tmp_path):
    case, *_, approval = approved_assessment(session, tmp_path)
    generate_markdown_report(session, case.case_id)
    revoke_approval(session, case.case_id, approval.approval_id, reviewer="Synthetic founder", reason="Re-review required.")
    assert not release_context(session, case.case_id)["approved"]
    assert "Readiness label: **CLIENT_READY**" not in generate_markdown_report(session, case.case_id).read_text(encoding="utf-8")


def test_evidence_review_rejects_wrong_entity_and_expired_records(session, tmp_path):
    case, _, _, row, _, proof = reviewed_assessment(session, tmp_path)
    for entity, expiry in (("Other company", "2099-01-01"), (case.organization, "2020-01-01")):
        with pytest.raises(ValueError):
            record_evidence_review(session, case.case_id, row.compliance_id, reviewer="Reviewer", notes="Controlled negative check", legal_entity=entity, file_path=proof, valid_until=expiry)


def test_source_bytes_changed_after_review_require_new_version(session, tmp_path):
    case, source, *_ = approved_assessment(session, tmp_path)
    Path(source.text_extract_path).write_text("Unregistered amendment", encoding="utf-8")
    assert not release_context(session, case.case_id)["approved"]
    with pytest.raises(ValueError):
        review_source(session, case.case_id, source.source_id, reviewer="Reviewer", notes="Attempt to review changed bytes in place.", requirements_complete=True)


def test_snapshot_and_archived_evidence_preserve_original_review(session, tmp_path):
    case, _, _, row, evidence, proof, approval = approved_assessment(session, tmp_path)
    original = proof.read_bytes()
    run = session.query(VerificationRun).filter_by(run_id=approval.verification_run_id).one()
    snapshot = json.loads(run.snapshot_json)
    assert snapshot["compliance_ledger"][0]["available_evidence"] == row.available_evidence
    proof.write_text("Later replacement", encoding="utf-8")
    row.available_evidence = "Later modified claim"
    session.flush()
    assert Path(evidence.archive_file_path).read_bytes() == original
    assert json.loads(run.snapshot_json) == snapshot
    assert not release_context(session, case.case_id)["approved"]


def test_natural_evidence_expiry_prevents_release_without_input_edit(session, tmp_path, monkeypatch):
    from tenderproof.services import decision_service

    case, _, _, _, evidence, *_ = approved_assessment(session, tmp_path)
    end = datetime.fromisoformat(evidence.valid_until).date()
    monkeypatch.setattr(decision_service, "today", lambda: end + timedelta(days=1))
    assert not release_context(session, case.case_id)["approved"]


def test_workbook_before_report_is_downloadable_draft(session, tmp_path):
    case, *_ = approved_assessment(session, tmp_path)
    export_case_to_excel(session, case.case_id)
    output = session.query(OutputPackage).filter_by(output_type="excel_export").one()
    assert output.readiness_label == "PRELIMINARY"
    assert output.approval_id is None


def test_approval_snapshot_survives_separate_sessions(session, tmp_path):
    from tenderproof.db import make_session

    case, *_ = reviewed_assessment(session, tmp_path)
    case_id = case.case_id
    session.commit()
    factory = make_session(session.bind.url.database)
    with factory() as check_session:
        run_verification(check_session, case_id)
        check_session.commit()
    with factory() as approval_session:
        approval = approve_case(approval_session, case_id, reviewer="Synthetic human", rationale="Separate-session review")
        approval_id = approval.approval_id
        approval_session.commit()
    with factory() as output_session:
        assert release_context(output_session, case_id)["approval"].approval_id == approval_id
        assert "Readiness label: **CLIENT_READY**" in generate_markdown_report(output_session, case_id).read_text(encoding="utf-8")
        output_session.commit()


def test_human_approval_ui_requires_confirmation_and_can_approve(session, tmp_path, monkeypatch):
    from streamlit.testing.v1 import AppTest
    from tenderproof.db import init_db, make_session

    case, *_ = reviewed_assessment(session, tmp_path)
    run_verification(session, case.case_id)
    case_id = case.case_id
    session.commit()
    monkeypatch.setenv("TENDERPROOF_DB_PATH", session.bind.url.database)
    init_db(session.bind.url.database)
    app = AppTest.from_file(str(Path(__file__).resolve().parents[1] / "app.py"))
    app.session_state["active_case_id"] = case_id
    app.run(timeout=40)
    app.sidebar.radio[0].set_value("Verification / Audit Log").run(timeout=40)
    next(w for w in app.text_input if w.label == "Approver name").set_value("Synthetic UI reviewer")
    next(w for w in app.text_area if w.label == "Approval rationale").set_value("Human reviewed this synthetic snapshot.")
    next(w for w in app.button if w.label == "Approve current assessment").click().run(timeout=40)
    assert not app.exception
    assert any("Confirm" in e.value for e in app.error)
    factory = make_session(session.bind.url.database)
    with factory() as checked:
        assert checked.query(DecisionApproval).count() == 0
    next(w for w in app.checkbox if w.label == "I approve the current evidence snapshot").set_value(True)
    next(w for w in app.button if w.label == "Approve current assessment").click().run(timeout=40)
    assert not app.exception
    assert not app.error
    with factory() as checked:
        assert checked.query(DecisionApproval).count() == 1


@pytest.mark.parametrize("content,expected", [(b"   ", "empty"), (b"valid prefix\xff\xff", "failed")])
def test_text_source_failures_are_explicit(tmp_path, content, expected):
    from tenderproof.services.source_service import extract_text

    path = tmp_path / "source.txt"
    path.write_bytes(content)
    _, status, _ = extract_text(path)
    assert status == expected


def test_manual_source_preserves_verbatim_original_bytes(session):
    text = "  Original clause\r\n with   spacing.\n"
    case = create_case(session, client_name="Synthetic source review")
    source = register_manual_text(session, case.case_id, text)
    assert Path(source.file_path).read_bytes() == text.encode("utf-8")


def test_modified_report_is_not_previewed_or_downloaded_as_approved(session, tmp_path, monkeypatch):
    from streamlit.testing.v1 import AppTest
    from tenderproof.db import init_db

    case, *_ = approved_assessment(session, tmp_path)
    report = generate_markdown_report(session, case.case_id)
    report.write_text("TAMPERED CLIENT_READY CONTENT", encoding="utf-8")
    case_id = case.case_id
    session.commit()
    monkeypatch.setenv("TENDERPROOF_DB_PATH", session.bind.url.database)
    init_db(session.bind.url.database)
    app = AppTest.from_file(str(Path(__file__).resolve().parents[1] / "app.py"))
    app.session_state["active_case_id"] = case_id
    app.run(timeout=40)
    app.sidebar.radio[0].set_value("Diagnostic Report Generator").run(timeout=40)
    assert not app.exception
    assert not any("TAMPERED CLIENT_READY" in item.value for item in app.markdown)
    downloads = app.get("download_button")
    assert downloads
    assert all(item.proto.disabled for item in downloads)
