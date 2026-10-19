"""Synthetic acceptance probes for the proposed safety contract, not implemented repairs."""

from uuid import uuid4

import pytest
from openpyxl import load_workbook
from sqlalchemy import create_engine
from sqlalchemy.orm import sessionmaker

from tenderproof.constants import REQUIRED_BID_CATEGORIES
from tenderproof.db import Base
from tenderproof.models import Case, ComplianceLedger, NormalizedRequirement, OutputPackage
from tenderproof.services import export_service, report_service, source_service
from tenderproof.services.bid_score_service import save_scores_and_recommendation
from tenderproof.services.case_service import update_case
from tenderproof.services.compliance_service import update_compliance_row
from tenderproof.services.egp_service import ensure_egp_checklist
from tenderproof.services.extraction_service import update_requirement
from tenderproof.services.verification_service import run_verification
from tenderproof.services.workflow_service import build_workflow_steps

READY = {"CLIENT_READY", "EXPORT_READY"}
POSITIVE = {"Pursue", "Pursue with partner", "Pursue only as sub-consultant"}


@pytest.fixture
def session(tmp_path, monkeypatch):
    for module, field in ((source_service, "EXTRACTED_TEXT_DIR"), (source_service, "UPLOADS_DIR"), (report_service, "REPORTS_DIR"), (export_service, "EXPORTS_DIR")):
        monkeypatch.setattr(module, field, tmp_path)
    engine = create_engine(f"sqlite:///{tmp_path / 'audit.db'}")
    Base.metadata.create_all(engine)
    with sessionmaker(bind=engine, autoflush=False, expire_on_commit=False)() as value:
        yield value
        value.rollback()
    engine.dispose()


def assessment(session, tmp_path, status="Met", reviewed=True, evidence="Synthetic license evidence reviewed by operator"):
    key = uuid4().hex[:10]
    case = Case(case_id=f"AUD-{key}", client_name="Synthetic audit client", tender_title="Synthetic tender")
    session.add(case)
    session.flush()
    source = source_service.register_manual_text(session, case.case_id, "Bidder must hold a current license.")
    req = NormalizedRequirement(requirement_id=f"REQ-{key}", case_id=case.case_id, source_id=source.source_id, tor_section_page="1", requirement_text="Bidder must hold a current license.", atomic_requirement_text="Bidder must hold a current license.", mandatory=True, needs_review=not reviewed, match_key=f"MATCH-{key}", canonical_anchor=f"{source.source_id}:1:1")
    row = ComplianceLedger(compliance_id=f"COMP-{key}", requirement_id=req.requirement_id, case_id=case.case_id, status=status, available_evidence=evidence)
    session.add_all([req, row])
    session.flush()
    scores = {c: {"score": 5, "rationale": "Synthetic opportunity fit", "evidence_or_gap_reference": req.requirement_id} for c in REQUIRED_BID_CATEGORIES}
    rec = save_scores_and_recommendation(session, case.case_id, scores)
    path = tmp_path / f"{key}.md"
    path.write_text("Synthetic diagnostic report for controlled audit.", encoding="utf-8")
    session.add(OutputPackage(output_id=f"OUT-{key}", case_id=case.case_id, output_type="markdown_report", file_path=str(path), readiness_label="DRAFT_ONLY", verification_status="not_run"))
    session.flush()
    return case, source, req, row, rec


@pytest.mark.parametrize("status", ["Gap", "Pending Input", "Partially Met", "Not Applicable"])
def test_mandatory_unresolved_blocks_positive_recommendation(session, tmp_path, status):
    *_, rec = assessment(session, tmp_path, status=status, evidence="")
    assert rec.recommendation not in POSITIVE


def test_unreviewed_requirement_blocks_readiness(session, tmp_path):
    case, *_ = assessment(session, tmp_path, reviewed=False)
    assert run_verification(session, case.case_id).readiness_label not in READY


def test_placeholder_evidence_blocks_readiness(session, tmp_path):
    case, *_ = assessment(session, tmp_path, evidence="Pending Input")
    assert run_verification(session, case.case_id).readiness_label not in READY


@pytest.mark.parametrize("evidence", ["License expired on 2020-01-01", "License belongs to another legal entity, not this bidder"])
def test_invalid_evidence_blocks_readiness(session, tmp_path, evidence):
    case, *_ = assessment(session, tmp_path, evidence=evidence)
    assert run_verification(session, case.case_id).readiness_label not in READY


@pytest.mark.parametrize("extraction_status", ["failed", "empty", "partial"])
def test_incomplete_source_extraction_blocks_readiness(session, tmp_path, extraction_status):
    case, source, *_ = assessment(session, tmp_path)
    source.extraction_status = extraction_status
    session.flush()
    assert run_verification(session, case.case_id).readiness_label not in READY


def test_new_source_invalidates_readiness(session, tmp_path):
    case, *_ = assessment(session, tmp_path)
    run_verification(session, case.case_id)
    source_service.register_manual_text(session, case.case_id, "Amendment: minimum turnover is now BDT 30000000.")
    assert case.readiness_label not in READY


def test_unreviewed_amendment_cannot_reclear_old_assessment(session, tmp_path):
    case, *_ = assessment(session, tmp_path)
    run_verification(session, case.case_id)
    source_service.register_manual_text(session, case.case_id, "Amendment: minimum turnover is now BDT 30000000.")
    assert run_verification(session, case.case_id).readiness_label not in READY


def test_compliance_change_invalidates_readiness(session, tmp_path):
    case, _, _, row, _ = assessment(session, tmp_path)
    run_verification(session, case.case_id)
    update_compliance_row(session, row.compliance_id, status="Gap", available_evidence="")
    assert case.readiness_label not in READY


def test_requirement_change_invalidates_readiness(session, tmp_path):
    case, _, req, *_ = assessment(session, tmp_path)
    run_verification(session, case.case_id)
    update_requirement(session, req.requirement_id, requirement_text="Minimum turnover is BDT 30000000.", atomic_requirement_text="Minimum turnover is BDT 30000000.", needs_review=True)
    assert case.readiness_label not in READY


def test_pending_egp_checklist_blocks_readiness(session, tmp_path):
    case, *_ = assessment(session, tmp_path)
    ensure_egp_checklist(session, case.case_id)
    assert run_verification(session, case.case_id).readiness_label not in READY


def test_warning_requires_preliminary_label(session, tmp_path):
    case, source, *_ = assessment(session, tmp_path)
    key = uuid4().hex[:10]
    req = NormalizedRequirement(requirement_id=f"REQ-{key}", case_id=case.case_id, source_id=source.source_id, requirement_text="Optional supporting material", atomic_requirement_text="Optional supporting material", mandatory=False, needs_review=False, match_key=key, canonical_anchor=key)
    session.add_all([req, ComplianceLedger(compliance_id=f"COMP-{key}", requirement_id=req.requirement_id, case_id=case.case_id, status="Gap", available_evidence="")])
    session.flush()
    summary = run_verification(session, case.case_id)
    assert summary.warnings, "Fixture must exercise a warning"
    assert summary.readiness_label == "PRELIMINARY"


def test_failed_verification_not_counted_as_complete_workflow(session, tmp_path):
    case, *_ = assessment(session, tmp_path, status="Met", evidence="")
    summary = run_verification(session, case.case_id)
    assert summary.overall_status == "failed"
    steps = build_workflow_steps(session, case)
    assert not next(s for s in steps if s["step"] == "Run verification")["complete"]


def test_human_approval_required_before_client_ready(session, tmp_path):
    case, *_ = assessment(session, tmp_path)
    assert run_verification(session, case.case_id).readiness_label not in READY


def test_stale_report_is_draft_after_evidence_removed(session, tmp_path):
    case, _, _, row, _ = assessment(session, tmp_path)
    run_verification(session, case.case_id)
    update_compliance_row(session, row.compliance_id, status="Gap", available_evidence="")
    text = report_service.generate_markdown_report(session, case.case_id).read_text(encoding="utf-8")
    assert "Readiness label: **CLIENT_READY**" not in text


def test_excel_does_not_execute_untrusted_formula_text(session, tmp_path):
    case, *_ = assessment(session, tmp_path)
    case.client_name = "=1+1"
    session.flush()
    path = export_service.export_case_to_excel(session, case.case_id)
    workbook = load_workbook(path, data_only=False)
    sheet = workbook["Case Summary"]
    col = next(c.column for c in sheet[1] if c.value == "client_name")
    assert sheet.cell(2, col).data_type != "f"


def test_case_update_cannot_self_assign_release_clearance(session, tmp_path):
    case, *_ = assessment(session, tmp_path, status="Gap", evidence="")
    update_case(session, case.case_id, readiness_label="EXPORT_READY", status="packaged")
    assert case.readiness_label not in READY


def test_control_synthetic_verified_row_can_be_scored(session, tmp_path):
    *_, rec = assessment(session, tmp_path)
    assert rec.recommendation == "Pursue"


def test_control_unknown_case_is_blocked(session):
    summary = run_verification(session, "NONEXISTENT-SYNTHETIC")
    assert summary.overall_status == "failed"
    assert summary.readiness_label == "BLOCKED"
